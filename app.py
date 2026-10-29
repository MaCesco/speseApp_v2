import base64, hashlib, io, json, os, re, secrets
from datetime import date, datetime, time, timedelta, timezone
from functools import wraps
import click
from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, abort, send_file
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

app = Flask(__name__)
uri = os.environ.get("DATABASE_URL", "sqlite:///spese.db")
if uri.startswith("postgres"):  # accetta qualsiasi forma della stringa di Neon
    uri = "postgresql+psycopg://" + uri.split("://", 1)[1]
app.config["SQLALCHEMY_DATABASE_URI"] = uri
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = "DATABASE_URL" in os.environ  # solo online (HTTPS)
app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {"pool_pre_ping": True, "pool_recycle": 280}  # database che va in pausa
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=14)  # la sessione scade dopo 14 giorni
_sk = os.environ.get("SECRET_KEY")
if not _sk:
    if "DATABASE_URL" in os.environ:  # online: senza chiave segreta le sessioni sarebbero falsificabili
        raise RuntimeError("Imposta la variabile d'ambiente SECRET_KEY (una stringa casuale lunga).")
    _sk = "solo-sviluppo-locale"
app.secret_key = _sk
db = SQLAlchemy(app)

ENTRATE = ["Stipendio", "Ripetizioni", "Concerti/Matrimoni", "Entrate extra"]
USCITE = ["Benzina auto", "Benzina moto", "Ricariche", "Biglietti treno/Autobus", "Parcheggio auto / moto",
          "Autostrada", "Libri", "Lavaggio Auto", "Abbigliamento", "Parrucchiera/Barbiere", "Supermercato",
          "Bar/Aperitivi", "Pranzi/Cene", "Cinema", "Discoteca", "Mostre/Fiere/Concerti/Giostre", "Viaggi",
          "Accessori/Computer", "Sport", "Assicurazione auto / moto", "Tagliandi auto / moto", "Accessori moto",
          "Farmaci / medicinali / Visite mediche", "Regali", "Corsi", "Uscite extra"]


class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(50), unique=True, nullable=False)
    password = db.Column(db.String(255), nullable=False)


class Movimento(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    data = db.Column(db.Date, nullable=False)
    tipo = db.Column(db.String(10), nullable=False)  # entrata | uscita
    categoria = db.Column(db.String(80), nullable=False)
    importo = db.Column(db.Float, nullable=False)
    descrizione = db.Column(db.String(200), default="")
    fonte = db.Column(db.String(10), default="manuale")  # manuale | import | ricorrente
    file_id = db.Column(db.Integer)  # file Excel (FileBilancio) da cui arriva, se importato


class Budget(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    categoria = db.Column(db.String(80), nullable=False)
    importo = db.Column(db.Float, nullable=False)


class Ricorrente(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    tipo = db.Column(db.String(10), nullable=False)
    categoria = db.Column(db.String(80), nullable=False)
    descrizione = db.Column(db.String(200), default="")
    importo = db.Column(db.Float, nullable=False)
    giorno = db.Column(db.Integer, nullable=False)
    ultimo = db.Column(db.String(7), nullable=False)  # ultimo mese generato, YYYY-MM


class Studente(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    nome = db.Column(db.String(120), nullable=False)
    tariffa = db.Column(db.Float, default=20)
    note = db.Column(db.String(500), default="")
    credito = db.Column(db.Float, default=0)  # euro pagati in anticipo, non ancora usati


class Lezione(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    studente_id = db.Column(db.Integer, db.ForeignKey("studente.id"), nullable=False, index=True)
    data = db.Column(db.Date, nullable=False)
    ora_inizio = db.Column(db.String(5), default="")
    ora_fine = db.Column(db.String(5), default="")
    ore = db.Column(db.Float, nullable=False)
    materia = db.Column(db.String(80), default="")
    modalita = db.Column(db.String(10), default="")  # Presenza | Online
    pagata = db.Column(db.Boolean, default=False, nullable=False)
    data_pagamento = db.Column(db.Date)
    fonte = db.Column(db.String(10), default="manuale")
    tariffa = db.Column(db.Float)  # tariffa applicata a questa lezione
    studente = db.relationship("Studente", backref="lezioni")


class FileBilancio(db.Model):
    """Elenco dei file Excel del bilancio importati (solo nome e data: i movimenti stanno in Movimento)."""
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    nome = db.Column(db.String(200), nullable=False)
    caricato_il = db.Column(db.DateTime, nullable=False)  # UTC
    contenuto = db.Column(db.LargeBinary)  # il file originale, cosi' si puo' riscaricare (vuoto per i file caricati prima)


class FileCaricato(db.Model):
    """Ultimo file Excel caricato (uno per tipo): resta salvato nel database online."""
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    tipo = db.Column(db.String(20), nullable=False)
    nome = db.Column(db.String(200), nullable=False)
    caricato_il = db.Column(db.DateTime, nullable=False)  # UTC
    contenuto = db.Column(db.LargeBinary, nullable=False)


class Stipendio(db.Model):
    """Una riga di Check_stipendi: i dati di un cedolino (mensile o tredicesima)."""
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    anno = db.Column(db.Integer, nullable=False)
    mese = db.Column(db.Integer, nullable=False)
    tipo = db.Column(db.String(10), nullable=False, default="mensile")  # mensile | tredicesima
    fonte = db.Column(db.String(10), default="sito")  # import | sito
    note = db.Column(db.String(1000))
    pdf = db.Column(db.String(200))


STIP_CAMPI = [  # (colonna Excel, attributo, titolo)
    ("B", "giorni", "Giorni lavorati"), ("C", "gg_mensa", "Giorni (trattenuta mensa)"), ("D", "scatti", "Scatti"),
    ("E", "unitario", "Compenso unitario"), ("F", "minimo", "Minimo"), ("G", "sup_ass", "Sup. ass."),
    ("H", "tot_retr", "Tot. retributivo"), ("I", "competenze", "Tot. competenze"), ("J", "trattenute", "Trattenute"),
    ("L", "fe_prec", "Ferie: anno prec."), ("M", "fe_spett", "Ferie: spettanti"), ("N", "fe_god", "Ferie: godute"),
    ("P", "par_prec", "PAR: anno prec."), ("Q", "par_spett", "PAR: spettanti"), ("R", "par_god", "PAR: godute"),
    ("U", "edp", "EDP"), ("V", "anticipo", "Anticipo for the better"), ("W", "mensa_tratt", "Trattenuta mensa"),
    ("X", "mensa_sup", "Supplemento mensa"), ("Z", "cometa", "Cometa"), ("AA", "tfr_fondo", "TFR fondo"), ("AB", "tfe", "TFE 12 AP")]
for _c, _a, _t in STIP_CAMPI:
    setattr(Stipendio, _a, db.Column(db.Float))


class CedolinoPdf(db.Model):
    """Il PDF di un cedolino, con i valori letti dal PDF (JSON)."""
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    anno = db.Column(db.Integer, nullable=False)
    mese = db.Column(db.Integer, nullable=False)
    tipo = db.Column(db.String(10), nullable=False, default="mensile")
    nome = db.Column(db.String(200), nullable=False)
    caricato_il = db.Column(db.DateTime, nullable=False)  # UTC
    contenuto = db.Column(db.LargeBinary, nullable=False)
    dati = db.Column(db.Text)  # {"v": valori, "voci": [...], "c": controlli}


ORE_CAMPI = ["lavoro", "sw", "ferie", "permessi", "fest_non", "fest_god", "trasferta", "pranzo_importo", "caffe", "spese"]


class OreGiorno(db.Model):
    """Un giorno di lavoro: ore (sede, smart working, ferie, ...) e piccole spese (pranzo, caffe')."""
    __table_args__ = (db.UniqueConstraint("user_id", "data", name="uq_ore_user_data"),)
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    data = db.Column(db.Date, nullable=False)
    pranzo = db.Column(db.String(10))  # None | mensa | fuori
    pranzo_nota = db.Column(db.String(200))
    nota = db.Column(db.String(300))


for _a in ORE_CAMPI:
    setattr(OreGiorno, _a, db.Column(db.Float))


class Impegno(db.Model):
    """Impegno extra dell'agenda (visita medica, concerto, ...)."""
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    data = db.Column(db.Date, nullable=False, index=True)
    ora_inizio = db.Column(db.String(5), default="")
    ora_fine = db.Column(db.String(5), default="")
    titolo = db.Column(db.String(150), nullable=False)
    luogo = db.Column(db.String(150), default="")
    nota = db.Column(db.String(500), default="")
    colore = db.Column(db.String(10), default="")  # "" = colore della sezione


class Scadenza(db.Model):
    """Un pagamento con validita' (assicurazione, bollo, ...). L'ultimo di ogni voce dice quando scade."""
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    voce = db.Column(db.String(120), nullable=False)
    importo = db.Column(db.Float)
    data_pagamento = db.Column(db.Date)
    validita_mesi = db.Column(db.Integer)
    scadenza = db.Column(db.Date, index=True)
    preavviso = db.Column(db.Integer, default=30)
    nota = db.Column(db.String(500), default="")
    chiusa = db.Column(db.Boolean, default=False, nullable=False)  # voce non piu' da rinnovare
    tipo = db.Column(db.String(10), default="scad")  # scad = ha una scadenza | nota = solo nota/pagamento
    colore = db.Column(db.String(10), default="")  # colore di sfondo scelto ("" = nessuno)


class Documento(db.Model):
    """Carte e documenti personali (identita', passaporto, patente, bancomat, carte, conti correnti)."""
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False, index=True)
    tipo = db.Column(db.String(12), default="altro")  # ci | pass | pat | nau | ts | banc | cc | conto | altro
    titolo = db.Column(db.String(120), nullable=False)
    intestatario = db.Column(db.String(120), default="")
    nome_p = db.Column(db.String(60), default="")
    cognome_p = db.Column(db.String(60), default="")
    numero = db.Column(db.String(60), default="")
    ente = db.Column(db.String(120), default="")  # ente di rilascio / banca
    rilascio = db.Column(db.Date)
    scadenza = db.Column(db.Date, index=True)
    iban = db.Column(db.String(40), default="")
    nota = db.Column(db.String(500), default="")
    colore = db.Column(db.String(10), default="")
    conto_id = db.Column(db.Integer)  # per le carte: id del conto associato
    immagine = db.Column(db.Text)  # immagine caricata (base64), solo per le carte
    img_tipo = db.Column(db.String(20))
    ordine = db.Column(db.Integer)  # posizione scelta dall'utente



def migra():
    """Aggiunge le colonne nuove alle tabelle gia' esistenti (database online)."""
    from sqlalchemy import inspect, text
    insp = inspect(db.engine)
    for tab, col, tipo in (("studente", "note", "VARCHAR(500)"), ("lezione", "tariffa", "FLOAT"), ("studente", "credito", "FLOAT"), ("movimento", "file_id", "INTEGER"), ("impegno", "colore", "VARCHAR(10)"), ("scadenza", "tipo", "VARCHAR(10)"), ("scadenza", "colore", "VARCHAR(10)"), ("documento", "conto_id", "INTEGER"), ("documento", "immagine", "TEXT"), ("documento", "img_tipo", "VARCHAR(20)"), ("documento", "ordine", "INTEGER"), ("documento", "nome_p", "VARCHAR(60)"), ("documento", "cognome_p", "VARCHAR(60)")):
        if col not in [c["name"] for c in insp.get_columns(tab)]:
            db.session.execute(text(f"ALTER TABLE {tab} ADD COLUMN {col} {tipo}"))
    if "contenuto" not in [c["name"] for c in insp.get_columns("file_bilancio")]:
        db.session.execute(text("ALTER TABLE file_bilancio ADD COLUMN contenuto " + ("BYTEA" if db.engine.dialect.name == "postgresql" else "BLOB")))
    db.session.execute(text("UPDATE lezione SET tariffa = (SELECT studente.tariffa FROM studente WHERE studente.id = lezione.studente_id) WHERE tariffa IS NULL"))
    db.session.execute(text("UPDATE studente SET credito = 0 WHERE credito IS NULL"))
    db.session.commit()


with app.app_context():
    db.create_all()
    migra()


# ---------- sicurezza ----------
def login_required(f):
    @wraps(f)
    def w(*a, **k):
        if "uid" not in session:
            return redirect(url_for("login"))
        return f(*a, **k)
    return w


@app.before_request
def csrf():
    if request.method == "POST":
        atteso = session.get("_t")
        if not atteso or not secrets.compare_digest(request.form.get("_t", ""), atteso):
            abort(400)


@app.context_processor
def ctx():
    session.setdefault("_t", secrets.token_hex(16))
    n_alert = 0
    n_doc = 0
    if "uid" in session:
        try:
            n_alert = len(alert_scadenze(session["uid"]))
            n_doc = len(documenti_in_scadenza(session["uid"]))
        except Exception:
            db.session.rollback()
    return {"csrf": session["_t"], "ENTRATE": ENTRATE, "USCITE": USCITE, "n_alert": n_alert, "n_doc": n_doc}


@app.cli.command("crea-utente")
@click.argument("username")
@click.password_option()
def crea_utente(username, password):
    if User.query.first():
        print("Esiste gia' un utente. Per cambiare password usa: flask --app app reimposta-password")
        return
    if len(password) < 10:
        print("Password troppo corta (minimo 10 caratteri).")
        return
    db.session.add(User(username=username, password=generate_password_hash(password)))
    db.session.commit()
    print("Utente creato.")


@app.cli.command("reimposta-password")
@click.password_option()
def reimposta_password(password):
    u = User.query.first()
    if not u:
        print("Nessun utente: usa prima crea-utente.")
    elif len(password) < 10:
        print("Password troppo corta (minimo 10 caratteri).")
    else:
        u.password = generate_password_hash(password)
        db.session.commit()
        print(f"Password aggiornata per l'utente: {u.username}")


# ---------- utilita ----------
def num(s):
    return float(str(s).replace(",", "."))


def nota_cella(c):
    """Testo della nota (commento) attaccata a una cella Excel, ripulito; stringa vuota se non c'e'."""
    if not c.comment:
        return ""
    t = (c.comment.text or "").strip()
    m = re.search(r"Comment:\s*\n(.*)", t, re.S)  # commenti "a catena" dell'Excel moderno
    if m:
        t = m.group(1).strip()
    else:  # nota classica: la prima riga e' spesso solo il nome dell'autore ("Marco:")
        righe = t.split("\n")
        if len(righe) > 1 and righe[0].rstrip().endswith(":") and len(righe[0]) <= 40:
            t = "\n".join(righe[1:]).strip()
    return " ".join(t.split())[:200]  # il campo descrizione accetta al massimo 200 caratteri


def applica_ricorrenti(uid):
    oggi = date.today()
    for r in Ricorrente.query.filter_by(user_id=uid):
        y, m = map(int, r.ultimo.split("-"))
        while True:
            m += 1
            if m > 12:
                y, m = y + 1, 1
            d = date(y, m, min(r.giorno, 28))
            if d > oggi:
                break
            db.session.add(Movimento(user_id=uid, data=d, tipo=r.tipo, categoria=r.categoria,
                                     importo=r.importo, descrizione=r.descrizione, fonte="ricorrente"))
            r.ultimo = f"{y}-{m:02d}"
    db.session.commit()


# ---------- pagine ----------
TENTATIVI = {}  # ip -> [istanti dei tentativi falliti]; in memoria, si azzera al riavvio
MAX_TENTATIVI, FINESTRA = 5, 15 * 60  # 5 errori ogni 15 minuti


def tentativi_recenti(ip):
    ora = datetime.now().timestamp()
    lista = [t for t in TENTATIVI.get(ip, []) if ora - t < FINESTRA]
    if lista:
        TENTATIVI[ip] = lista
    else:
        TENTATIVI.pop(ip, None)
    return lista


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        ip = request.remote_addr or "?"
        if len(tentativi_recenti(ip)) >= MAX_TENTATIVI:
            flash("Troppi tentativi falliti. Riprova tra 15 minuti.")
            return render_template("login.html"), 429
        u = User.query.filter_by(username=request.form.get("username", "").strip()).first()
        if u and check_password_hash(u.password, request.form.get("password", "")):
            TENTATIVI.pop(ip, None)
            session.clear()
            session.permanent = True
            session["uid"] = u.id
            return redirect(url_for("home"))
        TENTATIVI.setdefault(ip, []).append(datetime.now().timestamp())
        flash("Username o password errati.")
    return render_template("login.html")


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
@login_required
def home():
    return render_template("home.html", alert=alert_scadenze(session["uid"]), docs_alert=documenti_in_scadenza(session["uid"]), oggi=date.today())


@app.route("/dashboard")
@login_required
def dashboard():
    applica_ricorrenti(session["uid"])
    return render_template("dashboard.html")


@app.route("/api/dati")
@login_required
def dati():
    uid = session["uid"]
    mesi = {}
    q = db.session.query(Movimento.data, Movimento.tipo, Movimento.categoria, Movimento.importo)
    for d, t, c, i in q.filter(Movimento.user_id == uid):
        k = mesi.setdefault(f"{d:%Y-%m}", {"entrata": {}, "uscita": {}})[t]
        k[c] = round(k.get(c, 0) + i, 2)
    budget = {b.categoria: b.importo for b in Budget.query.filter_by(user_id=uid)}
    return jsonify(mesi=mesi, budget=budget)


@app.route("/nuovo", methods=["GET", "POST"])
@login_required
def nuovo():
    if request.method == "POST":
        f = request.form
        try:
            imp = abs(num(f["importo"]))
            d = datetime.strptime(f["data"], "%Y-%m-%d").date()
            cat = f["categoria"].strip()
            if not cat:
                raise ValueError
        except (ValueError, KeyError):
            flash("Controlla importo, data e categoria.")
            return redirect(url_for("nuovo"))
        db.session.add(Movimento(user_id=session["uid"], data=d, tipo="entrata" if f.get("tipo") == "entrata" else "uscita",
                                 categoria=cat, importo=imp, descrizione=f.get("descrizione", "").strip()))
        db.session.commit()
        flash("Movimento salvato.")
        return redirect(url_for("nuovo"))
    return render_template("nuovo.html", oggi=date.today().isoformat())


@app.route("/movimenti/<int:id>/modifica", methods=["GET", "POST"])
@login_required
def modifica_movimento(id):
    m = db.get_or_404(Movimento, id)
    if m.user_id != session["uid"]:
        abort(404)
    if request.method == "POST":
        f = request.form
        try:
            imp = abs(num(f["importo"]))
            d = datetime.strptime(f["data"], "%Y-%m-%d").date()
            cat = f["categoria"].strip()
            if not cat:
                raise ValueError
        except (ValueError, KeyError):
            flash("Controlla importo, data e categoria.")
            return redirect(url_for("modifica_movimento", id=id))
        m.data, m.importo, m.categoria = d, imp, cat
        m.tipo = "entrata" if f.get("tipo") == "entrata" else "uscita"
        m.descrizione = f.get("descrizione", "").strip()
        db.session.commit()
        flash("Movimento aggiornato.")
        return redirect(url_for("movimenti", mese=f"{d:%Y-%m}"))
    return render_template("nuovo.html", oggi=m.data.isoformat(), m=m)


@app.route("/movimenti")
@login_required
def movimenti():
    mese = request.args.get("mese", "")
    try:
        y, m = map(int, mese.split("-"))
        a = date(y, m, 1)
    except ValueError:
        a = date.today().replace(day=1)
    b = date(a.year + (a.month == 12), a.month % 12 + 1, 1)
    righe = (Movimento.query.filter(Movimento.user_id == session["uid"], Movimento.data >= a, Movimento.data < b)
             .order_by(Movimento.data.desc(), Movimento.id.desc()).all())
    return render_template("movimenti.html", righe=righe, mese=f"{a:%Y-%m}")


@app.route("/elimina/<int:id>", methods=["POST"])
@login_required
def elimina(id):
    r = db.get_or_404(Movimento, id)
    if r.user_id != session["uid"]:
        abort(404)
    mese = f"{r.data:%Y-%m}"
    db.session.delete(r)
    db.session.commit()
    return redirect(url_for("movimenti", mese=mese))


@app.route("/impostazioni", methods=["GET", "POST"])
@login_required
def impostazioni():
    uid = session["uid"]
    if request.method == "POST":
        f = request.form
        az = f.get("azione")
        try:
            if az == "budget":
                b = Budget.query.filter_by(user_id=uid, categoria=f["categoria"]).first() or Budget(user_id=uid, categoria=f["categoria"])
                b.importo = num(f["importo"])
                db.session.add(b)
            elif az == "ricorrente":
                t = date.today()
                y, m = (t.year, t.month - 1) if t.month > 1 else (t.year - 1, 12)
                db.session.add(Ricorrente(user_id=uid, tipo="entrata" if f["tipo"] == "entrata" else "uscita",
                                          categoria=f["categoria"].strip(), descrizione=f.get("descrizione", "").strip(),
                                          importo=abs(num(f["importo"])), giorno=min(max(int(f["giorno"]), 1), 28),
                                          ultimo=f"{y}-{m:02d}"))
            elif az == "del_budget":
                Budget.query.filter_by(id=int(f["id"]), user_id=uid).delete()
            elif az == "del_ric":
                Ricorrente.query.filter_by(id=int(f["id"]), user_id=uid).delete()
            db.session.commit()
        except (ValueError, KeyError):
            db.session.rollback()
            flash("Dati non validi.")
        return redirect(url_for("impostazioni"))
    file_bilanci, anni_senza_file = elenco_bilanci(uid)
    return render_template("impostazioni.html", budget=Budget.query.filter_by(user_id=uid).all(),
                           ric=Ricorrente.query.filter_by(user_id=uid).all(),
                           file_bilanci=file_bilanci, anni_senza_file=anni_senza_file,
                           anni_export=sorted({int(y) for (y,) in db.session.query(db.extract("year", Movimento.data)).filter(Movimento.user_id == uid).distinct()}))


@app.route("/password", methods=["POST"])
@login_required
def cambia_password():
    u = db.session.get(User, session["uid"])
    f = request.form
    if not check_password_hash(u.password, f.get("vecchia", "")):
        flash("La password attuale non e' corretta.")
    elif len(f.get("nuova", "")) < 10 or f["nuova"] != f.get("conferma"):
        flash("La nuova password deve avere almeno 10 caratteri ed essere uguale nei due campi.")
    else:
        u.password = generate_password_hash(f["nuova"])
        db.session.commit()
        flash("Password aggiornata.")
    return redirect(url_for("impostazioni"))


@app.route("/importa", methods=["POST"])
@login_required
def importa():
    uid = session["uid"]
    n, anni, errori = 0, set(), []
    for file in request.files.getlist("file"):
        raw = file.read()
        try:
            wb = load_workbook(io.BytesIO(raw), data_only=True)
        except Exception:
            errori.append(file.filename)
            continue
        nuovi = []
        for ws in wb:
            if ws.title.strip().upper() == "ANNUALE":
                continue
            for c in range(3, ws.max_column + 1):
                d = ws.cell(2, c).value
                if not isinstance(d, datetime):
                    continue
                for r in range(3, 33):
                    v, cat = ws.cell(r, c).value, ws.cell(r, 1).value
                    if cat and isinstance(v, (int, float)) and v:
                        cat = str(cat).strip()
                        nuovi.append(Movimento(user_id=uid, data=d.date(), tipo="entrata" if cat in ENTRATE else "uscita",
                                               categoria=cat, importo=abs(v), fonte="import",
                                               descrizione=nota_cella(ws.cell(r, c))))
        if not nuovi:
            errori.append(file.filename)
            continue
        ys = {m.data.year for m in nuovi}
        da_sito = (wb.properties.title or "") == "spese-export"  # scaricato da qui: contiene gia' anche i movimenti inseriti dal sito
        for y in ys:  # sostituisce solo gli anni presenti nel file
            q = Movimento.query.filter(Movimento.user_id == uid, Movimento.data >= date(y, 1, 1), Movimento.data < date(y + 1, 1, 1))
            (q if da_sito else q.filter(Movimento.fonte == "import")).delete()
        rec = FileBilancio(user_id=uid, nome=(file.filename or "Bilancio.xlsx")[:200],
                           caricato_il=datetime.now(timezone.utc).replace(tzinfo=None), contenuto=raw)
        db.session.add(rec)
        db.session.flush()
        for m in nuovi:
            m.file_id = rec.id
        db.session.add_all(nuovi)
        n += len(nuovi)
        anni |= ys
    db.session.flush()
    vuoti = [i for (i,) in db.session.query(FileBilancio.id).outerjoin(Movimento, Movimento.file_id == FileBilancio.id)
             .filter(FileBilancio.user_id == uid, Movimento.id.is_(None))]  # file i cui anni sono stati sostituiti da uno nuovo
    if vuoti:
        FileBilancio.query.filter(FileBilancio.id.in_(vuoti)).delete(synchronize_session=False)
    db.session.commit()
    if anni:
        flash(f"Importati {n} movimenti per gli anni: {', '.join(map(str, sorted(anni)))}.")
    if errori:
        flash("File non letti: " + ", ".join(errori))
    return redirect(url_for("impostazioni"))


@app.route("/importa/scarica/<int:id>")
@login_required
def importa_scarica(id):
    """Il file Excel esattamente come era stato caricato."""
    f = mio(FileBilancio, id)
    if not f.contenuto:
        abort(404)
    return send_file(io.BytesIO(f.contenuto), as_attachment=True, download_name=f.nome)


@app.route("/esporta/<int:anno>")
@login_required
def esporta_anno(anno):
    """Excel dell'anno con tutti i movimenti attuali (importati e inseriti dal sito), nello stesso formato dell'importazione:
    un foglio per mese, i giorni in colonna, le categorie in riga, le note come commenti della cella. Si puo' ricaricare."""
    if not 1900 <= anno <= 2200:
        abort(404)
    mov = (Movimento.query.filter(Movimento.user_id == session["uid"], Movimento.data >= date(anno, 1, 1), Movimento.data < date(anno + 1, 1, 1))
           .order_by(Movimento.data, Movimento.id).all())
    if not mov:
        flash(f"Non ci sono movimenti nel {anno} da esportare.")
        return redirect(url_for("impostazioni"))
    mesi = "Gennaio Febbraio Marzo Aprile Maggio Giugno Luglio Agosto Settembre Ottobre Novembre Dicembre".split()
    grassetto = Font(bold=True)
    ordine = {c: i for i, c in enumerate(ENTRATE + USCITE)}
    wb = Workbook()
    ann = wb.active
    ann.title = "ANNUALE"
    cats = sorted({m.categoria for m in mov}, key=lambda c: (c not in ENTRATE, ordine.get(c, 999), c))
    ann.append(["Categoria", "Totale"] + mesi)
    for c in ann[1]:
        c.font = grassetto
    for cat in cats:
        riga = [round(sum(m.importo for m in mov if m.categoria == cat and m.data.month == i), 2) or None for i in range(1, 13)]
        ann.append([cat, round(sum(m.importo for m in mov if m.categoria == cat), 2)] + riga)
    ann.column_dimensions["A"].width = 24
    for i in range(1, 13):
        ws = wb.create_sheet(mesi[i - 1])
        dm = [m for m in mov if m.data.month == i]
        ws["A1"] = f"{mesi[i - 1]} {anno}"
        ws["A1"].font = grassetto
        ws["A2"], ws["B2"] = "Categoria", "Totale"
        ws["A2"].font = ws["B2"].font = grassetto
        giorni = (date(anno + (i == 12), i % 12 + 1, 1) - date(anno, i, 1)).days
        for g in range(1, giorni + 1):
            c = ws.cell(2, 2 + g, datetime(anno, i, g))
            c.number_format = "D"
            c.font = grassetto
        cm = sorted({m.categoria for m in dm}, key=lambda c: (c not in ENTRATE, ordine.get(c, 999), c))
        for r, cat in enumerate(cm, 3):
            ws.cell(r, 1, cat)
            tot = 0
            for g in range(1, giorni + 1):
                qui = [m for m in dm if m.categoria == cat and m.data.day == g]
                if not qui:
                    continue
                v = round(sum(m.importo for m in qui), 2)
                tot += v
                cella = ws.cell(r, 2 + g, v)
                note = " · ".join(dict.fromkeys(m.descrizione.strip() for m in qui if m.descrizione and m.descrizione.strip()))
                if note:
                    cella.comment = Comment(note, "Le mie spese")
            ws.cell(r, 2, round(tot, 2))
        r = max(35, 4 + len(cm))  # i totali stanno sotto le righe delle categorie, cosi' non vengono letti come movimenti
        ws.cell(r, 1, "Totale entrate").font = grassetto
        ws.cell(r, 2, round(sum(m.importo for m in dm if m.tipo == "entrata"), 2))
        ws.cell(r + 1, 1, "Totale uscite").font = grassetto
        ws.cell(r + 1, 2, round(sum(m.importo for m in dm if m.tipo == "uscita"), 2))
        ws.column_dimensions["A"].width = 24
        ws.freeze_panes = "C3"
    wb.properties.title = "spese-export"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Bilancio_{anno}_aggiornato_{date.today()}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/importa/elimina/<int:id>", methods=["POST"])
@login_required
def importa_elimina(id):
    f = mio(FileBilancio, id)
    nome = f.nome
    n = Movimento.query.filter_by(user_id=session["uid"], file_id=id, fonte="import").delete()
    db.session.delete(f)
    db.session.commit()
    flash(f"File «{nome}» eliminato, con i suoi {n} movimenti. I movimenti inseriti a mano non sono stati toccati.")
    return redirect(url_for("impostazioni"))


@app.route("/importa/elimina-anno/<int:anno>", methods=["POST"])
@login_required
def importa_elimina_anno(anno):
    """Per i movimenti importati prima che i file venissero registrati (nessun file associato)."""
    if not 1900 <= anno <= 2200:
        abort(404)
    n = Movimento.query.filter(Movimento.user_id == session["uid"], Movimento.fonte == "import", Movimento.file_id.is_(None),
                               Movimento.data >= date(anno, 1, 1), Movimento.data < date(anno + 1, 1, 1)).delete()
    db.session.commit()
    flash(f"Eliminati {n} movimenti importati del {anno}. I movimenti inseriti a mano non sono stati toccati.")
    return redirect(url_for("impostazioni"))


def elenco_bilanci(uid):
    """File importati (con anni e numero di movimenti) + anni importati senza file associato."""
    anno = db.extract("year", Movimento.data)
    per_file = {}
    q = (db.session.query(Movimento.file_id, anno, db.func.count())
         .filter(Movimento.user_id == uid, Movimento.fonte == "import").group_by(Movimento.file_id, anno))
    for fid, y, c in q:
        per_file.setdefault(fid, {})[int(y)] = c
    con_originale = {i for (i,) in db.session.query(FileBilancio.id).filter(FileBilancio.user_id == uid, FileBilancio.contenuto.isnot(None))}
    file = [dict(f=f, anni=sorted(per_file.get(f.id, {})), n=sum(per_file.get(f.id, {}).values()),
                 iso=f.caricato_il.isoformat() + "Z", originale=f.id in con_originale)
            for f in FileBilancio.query.options(db.defer(FileBilancio.contenuto)).filter_by(user_id=uid)
            .order_by(FileBilancio.caricato_il.desc(), FileBilancio.id.desc())]
    vecchi = sorted(per_file.get(None, {}).items())
    return file, vecchi


# ---------- ripetizioni ----------
MESI_IT = {m: i for i, m in enumerate("gennaio febbraio marzo aprile maggio giugno luglio agosto settembre ottobre novembre dicembre".split(), 1)}
PAG = re.compile(r"pagat|bonifico|contanti|paypal|ricevut|incass|pagamento|regalat", re.I)


def mio(model, id):
    o = db.get_or_404(model, id)
    if o.user_id != session["uid"]:
        abort(404)
    return o


def parse_data(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, str):
        m = re.search(r"(\d{1,2})\s+([a-z]+)\s+(\d{4})", v.lower())
        if m and m.group(2) in MESI_IT:
            return date(int(m.group(3)), MESI_IT[m.group(2)], int(m.group(1)))
    return None


def rosso(c):
    """Il testo della cella e' rosso? (rosso = lezione da pagare, verde = pagata)"""
    col = c.font.color
    if col is None or col.type != "rgb" or not isinstance(col.rgb, str):
        return False
    h = col.rgb[-6:]
    return int(h[0:2], 16) > 180 and int(h[2:4], 16) < 90 and int(h[4:6], 16) < 90


def leggi_credito(ws):
    """Credito prepagato scritto dal sito nel foglio (colonna L). None se il foglio non lo contiene."""
    v = ws.cell(2, 12).value
    if str(ws.cell(1, 12).value or "").strip().lower() == "credito eur" and isinstance(v, (int, float)) and v >= 0:
        return float(v)
    return None


NOTA_PAG = re.compile(r"(\d+)\s*lezion[ei]\s*(?:gi[aà]\s*)?pagat", re.I)
NOTA_EURO = re.compile(r"(\d+(?:[.,]\d+)?)\s*(?:€|eur)", re.I)


def credito_da_note(ws, tariffa):
    """Nel tuo Excel un pagamento anticipato e' scritto come nota sulla prima lezione che copre, ad es. «4 lezioni pagate bonifico 16/10 - 100€».
    Se da quella riga in poi ci sono meno lezioni di quelle pagate, le altre sono gia' pagate e non ancora svolte: restano come credito."""
    righe = []  # (riga, numero di lezioni pagate scritte nella nota, importo nella nota)
    lez = []
    for r in range(2, ws.max_row + 1):
        d, ore = parse_data(ws.cell(r, 1).value), ws.cell(r, 4).value
        if not d or not isinstance(ore, (int, float)) or ore <= 0:
            continue
        lez.append(r)
        for c in range(5, 10):
            v = ws.cell(r, c).value
            m = NOTA_PAG.search(v) if isinstance(v, str) else None
            if m:
                e = NOTA_EURO.search(v)
                righe.append((len(lez) - 1, int(m.group(1)), float(e.group(1).replace(",", ".")) if e else None))
    if not righe:
        return 0.0
    i, n, euro = righe[-1]
    resto = n - (len(lez) - i)
    if resto <= 0 or n <= 0 or any(rosso(ws.cell(r, c)) for r in lez[i:] for c in range(1, 6)):
        return 0.0  # tutte svolte, oppure ci sono lezioni segnate da pagare: non indovino nulla
    return round((euro / n if euro else tariffa) * resto, 2)


def leggi_foglio(ws):
    tariffa = 20.0
    for r in range(1, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if isinstance(a, str) and a.strip().lower().startswith("importi"):
            m = re.search(r"\d+(?:[.,]\d+)?", str(b or ""))
            if m:
                tariffa = float(m.group().replace(",", "."))
            break
    out = []
    usa_j = str(ws.cell(1, 10).value or "").strip().lower().startswith("tariffa")

    def riga_tar(r):
        v = ws.cell(r, 10).value if usa_j else None
        return float(v) if isinstance(v, (int, float)) and v >= 0 else tariffa
    for r in range(2, ws.max_row + 1):
        d, ore = parse_data(ws.cell(r, 1).value), ws.cell(r, 4).value
        if not d or not isinstance(ore, (int, float)) or ore <= 0:
            continue
        note = [str(ws.cell(r, c).value) for c in range(5, 10) if isinstance(ws.cell(r, c).value, str)]
        testo = " ".join(note).lower()
        e = ws.cell(r, 5).value
        modalita = "Presenza" if "presenza" in testo else "Online" if re.search(r"dad|skype|meet", testo) else ""
        hh = lambda v: v.strftime("%H:%M") if isinstance(v, (time, datetime)) else ""
        out.append(dict(data=d, ora_inizio=hh(ws.cell(r, 2).value), ora_fine=hh(ws.cell(r, 3).value), ore=float(ore),
                        materia=(str(e).strip()[:80] if isinstance(e, str) else ""), modalita=modalita,
                        pagata=not any(rosso(ws.cell(r, c)) for c in range(1, 6)), tariffa=riga_tar(r)))
    nota = ""
    for r in (2, 3):
        for c in range(6, 10):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip() and not PAG.search(v) and not re.fullmatch(r"(in )?presenza|dad|skype|google meet|meet", v.strip(), re.I):
                nota = nota or v.strip()[:300]
    k2 = ws.cell(2, 11).value
    if str(ws.cell(1, 11).value or "").strip().lower() == "note" and isinstance(k2, str) and k2.strip():
        nota = k2.strip()[:500]
    return tariffa, nota, out


def txt_pre(k):
    return "nessuna lezione prepagata" if k == 0 else "1 lezione già pagata da svolgere" if k == 1 else f"{k} lezioni già pagate da svolgere"


def n_prepagate(s):
    """Quante lezioni intere copre il credito dell'allievo."""
    t = s.tariffa or 0
    return int((s.credito or 0) / t + 1e-9) if t > 0 else 0


@app.route("/ripetizioni")
@login_required
def ripetizioni():
    uid = session["uid"]
    al = {s.id: dict(s=s, n=0, dovuto=0.0, incassato=0.0, inc_anno=0.0, prima=None, ultima=None) for s in Studente.query.filter_by(user_id=uid)}
    inc_mesi, ult, oggi_d = {}, {}, date.today()
    q = db.session.query(Lezione.studente_id, Lezione.ore, Lezione.tariffa, Lezione.pagata, Lezione.data).filter(Lezione.user_id == uid)
    for sid, ore, tar, pag, d in q:
        a = al[sid]
        euro = ore * (tar or a["s"].tariffa or 0)
        kk = (d <= oggi_d, d)  # l'ultima lezione gia' svolta; se non ce ne sono, la piu' recente
        if sid not in ult or kk > ult[sid][0]:
            ult[sid] = (kk, tar if tar is not None else a["s"].tariffa)
        if pag:
            a["incassato"] += euro
            if d.year == oggi_d.year:
                a["inc_anno"] += euro
            k = d.strftime("%Y-%m")
            inc_mesi[k] = round(inc_mesi.get(k, 0) + euro, 2)
        else:
            a["n"] += 1
            a["dovuto"] += euro
        if not a["ultima"] or d > a["ultima"]:
            a["ultima"] = d
        if not a["prima"] or d < a["prima"]:
            a["prima"] = d
    # prima chi ha lezioni da pagare; poi i piu' recenti (prima lezione piu' vicina a oggi) in alto, i piu' storici in basso
    for a in al.values():
        a["pre"] = n_prepagate(a["s"])
    elenco = sorted(al.values(), key=lambda a: (a["n"] == 0, -(a["prima"] or date.max).toordinal(), -a["s"].id))
    anno = date.today() - timedelta(days=365)
    tutti = request.args.get("tutti") == "1"
    visibili = [a for a in elenco if tutti or a["n"] or (a["ultima"] and a["ultima"] >= anno)]
    tariffe = {a["s"].nome: (ult[a["s"].id][1] if a["s"].id in ult else a["s"].tariffa) or 0 for a in elenco}
    return render_template("ripetizioni.html", anno_corr=oggi_d.year, tariffe=tariffe, elenco=elenco, visibili=visibili, tutti=tutti, oggi=date.today().isoformat(),
                           inc_mesi=inc_mesi, tot=sum(a["dovuto"] for a in elenco), tot_n=sum(a["n"] for a in elenco),
                           n_deb=sum(1 for a in elenco if a["n"]), p=ore_periodi(uid), tf=tariffe_periodi(uid), ultimo_file=ultimo_file(uid))


@app.route("/ripetizioni/<int:id>")
@login_required
def allievo(id):
    s = mio(Studente, id)
    lez = Lezione.query.filter_by(studente_id=id).order_by(Lezione.data.desc(), Lezione.ora_inizio.desc()).all()
    n = sum(1 for l in lez if not l.pagata)
    oggi = date.today()
    credito = s.credito or 0
    val = lambda l: l.ore * (l.tariffa or s.tariffa or 0)
    yr = oggi.year
    svolte = [l for l in lez if l.data <= oggi]
    # fatturato (ore x tariffa, pagato o no) per mese e per settimana, dalle lezioni gia' svolte
    import calendar
    per_mese, per_sett, ore_mese = {}, {}, {}
    for l in svolte:
        ore_mese[(l.data.year, l.data.month)] = ore_mese.get((l.data.year, l.data.month), 0) + l.ore
        per_mese[(l.data.year, l.data.month)] = per_mese.get((l.data.year, l.data.month), 0) + val(l)
        lun = l.data - timedelta(days=l.data.weekday())
        per_sett[lun] = per_sett.get(lun, 0) + val(l)
    cur_m = (oggi.year, oggi.month)
    prec_m = (oggi.year, oggi.month - 1) if oggi.month > 1 else (oggi.year - 1, 12)
    ms = [k for k in per_mese if k != cur_m] or list(per_mese)
    somma_m = sum(per_mese[k] for k in ms)
    media_m = somma_m / len(ms) if ms else 0
    media_s = somma_m / (sum(calendar.monthrange(*k)[1] for k in ms) / 7) if ms else 0
    somma_h = sum(ore_mese[k] for k in ms)
    ore_m = somma_h / len(ms) if ms else 0
    ore_s = somma_h / (sum(calendar.monthrange(*k)[1] for k in ms) / 7) if ms else 0
    lun0 = oggi - timedelta(days=oggi.weekday())
    # allievo fermo da tempo: le medie sono vecchie, indichiamo a che periodo si riferiscono
    vecchia = ""
    if svolte and per_mese.get(cur_m, 0) == 0 and per_mese.get(prec_m, 0) == 0:
        ult = max(l.data for l in svolte)
        vecchia = f"al {ult.year}" if ult.year != oggi.year else f"a {MESI_LUNGHI[ult.month - 1]} {ult.year}"
    medie = dict(ore_m=ore_m, ore_s=ore_s, vecchia=vecchia, nome_mese=MESI_LUNGHI[oggi.month - 1], m=media_m, s=media_s, m_cur=per_mese.get(cur_m, 0), m_prec=per_mese.get(prec_m, 0),
                 s_cur=per_sett.get(lun0, 0), s_prec=per_sett.get(lun0 - timedelta(days=7), 0))
    return render_template("allievo.html", anno_corr=yr, medie=medie,
                           inc_a=sum(val(l) for l in lez if l.pagata and l.data.year == yr),
                           inc_p=sum(val(l) for l in lez if l.pagata and l.data.year == yr - 1),
                           ore_a=sum(l.ore for l in svolte if l.data.year == yr),
                           ore_p=sum(l.ore for l in svolte if l.data.year == yr - 1),
                           fatturato_tot=sum(val(l) for l in svolte),
                           s=s, lez=lez, n=n, n_tot=len(lez),
                           ore_svolte=sum(l.ore for l in lez if l.data <= oggi), ore_prog=sum(l.ore for l in lez if l.data > oggi),
                           dovuto=sum(l.ore * (l.tariffa or s.tariffa or 0) for l in lez if not l.pagata),
                           # totale incassato: lezioni pagate (anche future) + credito versato in anticipo
                           incassato=sum(l.ore * (l.tariffa or s.tariffa or 0) for l in lez if l.pagata) + credito,
                           debito_ore=sum(l.ore for l in lez if not l.pagata),
                           credito_ore=round(credito / s.tariffa, 2) if (s.tariffa or 0) > 0 else 0,
                           da=(request.args.get("da", "") if request.args.get("da", "").startswith("/agenda") else ""),
                           serie=[[l.data.isoformat(), l.ore, round(val(l), 2) if l.pagata else 0] for l in lez],
                           credito=credito, n_pre=n_prepagate(s))


@app.route("/ripetizioni/nuova", methods=["POST"])
@login_required
def rip_nuova():
    uid, f = session["uid"], request.form
    try:
        d = datetime.strptime(f["data"], "%Y-%m-%d").date()
        ore = round((datetime.strptime(f["fine"], "%H:%M") - datetime.strptime(f["inizio"], "%H:%M")).total_seconds() / 3600, 2)
        nome = f["studente"].strip()
        if ore <= 0 or not nome:
            raise ValueError
    except (ValueError, KeyError):
        flash("Controlla allievo, data e orari.")
        return redirect(url_for("ripetizioni"))
    s = Studente.query.filter_by(user_id=uid, nome=nome).first()
    try:  # l'importo scritto nel modulo vale sempre per questa lezione, anche per un allievo gia' presente
        t_in = abs(num(f["tariffa"])) if (f.get("tariffa") or "").strip() else None
    except ValueError:
        t_in = None
    if not s:
        s = Studente(user_id=uid, nome=nome, tariffa=t_in or 0)
        db.session.add(s)
        db.session.flush()
        t_lez = s.tariffa
    elif t_in is not None:
        t_lez = t_in
    else:  # campo vuoto: l'importo dell'ultima lezione svolta
        base = Lezione.query.filter_by(studente_id=s.id).filter(Lezione.tariffa.isnot(None))
        ordine = (Lezione.data.desc(), Lezione.ora_inizio.desc())
        u = base.filter(Lezione.data <= date.today()).order_by(*ordine).first() or base.order_by(*ordine).first()
        t_lez = u.tariffa if u else (s.tariffa or 0)
    l = Lezione(user_id=uid, studente_id=s.id, data=d, ora_inizio=f["inizio"], ora_fine=f["fine"], ore=ore,
                materia=f.get("materia", "").strip(), modalita=f.get("modalita", ""), tariffa=t_lez)
    costo = ore * (t_lez or 0)
    arretrati = Lezione.query.filter_by(studente_id=s.id, pagata=False).count()
    prepagata = arretrati == 0 and costo > 0 and (s.credito or 0) >= costo - 0.001
    if prepagata:  # l'allievo aveva gia' pagato in anticipo: la lezione nasce pagata e scala il credito
        l.pagata, l.data_pagamento = True, date.today()
        s.credito = round(max((s.credito or 0) - costo, 0), 2)
    db.session.add(l)
    db.session.commit()
    flash(f"Lezione aggiunta e già pagata con il credito. Restano: {txt_pre(n_prepagate(s))}." if prepagata else "Lezione aggiunta.")
    return redirect(url_for("allievo", id=s.id))


@app.route("/ripetizioni/<int:id>/dati", methods=["POST"])
@login_required
def rip_dati(id):
    s, f = mio(Studente, id), request.form
    try:
        nome, tar = f["nome"].strip(), abs(num(f["tariffa"]))
        if not nome:
            raise ValueError
    except (ValueError, KeyError):
        flash("Nome o tariffa non validi.")
        return redirect(url_for("allievo", id=id))
    s.nome, s.tariffa, s.note = nome, tar, f.get("note", "").strip()[:500]
    if "credito" in f:
        try:
            s.credito = abs(num(f["credito"] or 0))
        except ValueError:
            flash("Credito non valido: gli altri dati sono stati salvati.")
    if f.get("applica"):
        Lezione.query.filter_by(studente_id=id, pagata=False).update({"tariffa": tar})
    db.session.commit()
    flash("Dati salvati.")
    return redirect(url_for("allievo", id=id))


@app.route("/ripetizioni/<int:id>/elimina", methods=["POST"])
@login_required
def rip_elimina(id):
    s = mio(Studente, id)
    Lezione.query.filter_by(studente_id=id).delete()
    db.session.delete(s)
    db.session.commit()
    flash("Allievo eliminato.")
    return redirect(url_for("ripetizioni"))


@app.route("/ripetizioni/<int:id>/incassa", methods=["POST"])
@login_required
def rip_incassa(id):
    s, f = mio(Studente, id), request.form
    q = Lezione.query.filter_by(studente_id=id, pagata=False).order_by(Lezione.data, Lezione.id)
    n = 0
    try:
        residuo = None if f.get("tutte") else abs(num(f["importo"])) + (s.credito or 0)  # importo + credito gia' presente
    except (ValueError, KeyError):
        flash("Importo non valido.")
        return redirect(url_for("allievo", id=id))
    for l in q:
        costo = l.ore * (l.tariffa or s.tariffa or 0)
        if residuo is not None:
            if costo > residuo + 0.001:
                break
            residuo -= costo
        l.pagata, l.data_pagamento = True, date.today()
        n += 1
    msg = "1 lezione segnata come pagata." if n == 1 else f"{n} lezioni segnate come pagate."
    if residuo is not None:
        s.credito = round(max(residuo, 0), 2)  # quello che avanza e' un anticipo per le prossime lezioni
        if s.credito:
            k = n_prepagate(s)
            msg += f" Avanzano {s.credito:.0f} €: {txt_pre(k)}." if k else f" Avanzano {s.credito:.0f} € di credito."
    db.session.commit()
    flash(msg)
    return redirect(url_for("allievo", id=id))


@app.route("/lezione/<int:id>/<azione>", methods=["POST"])
@login_required
def rip_lezione(id, azione):
    l = mio(Lezione, id)
    sid = l.studente_id
    if azione == "paga":
        l.pagata = not l.pagata
        l.data_pagamento = date.today() if l.pagata else None
    elif azione == "elimina":
        db.session.delete(l)
    db.session.commit()
    return redirect(url_for("allievo", id=sid))


@app.route("/ripetizioni/importa", methods=["POST"])
@login_required
def rip_importa():
    uid = session["uid"]
    f = request.files.get("file")
    dati = f.read() if f else b""
    try:
        wb = load_workbook(io.BytesIO(dati), data_only=True)
        fogli = [(ws.title.strip(),) + leggi_foglio(ws) for ws in wb]  # (nome, tariffa, nota, lezioni)
        crediti = {ws.title.strip(): leggi_credito(ws) for ws in wb}
        da_note = {nome: credito_da_note(ws, tar) for ws, (nome, tar, _n, _l) in zip(wb, fogli)}  # solo se il foglio non ha la colonna Credito
    except Exception:
        flash("Il file non e' un Excel (.xlsx) valido: non ho cambiato nulla.")
        return redirect(url_for("ripetizioni"))
    fogli = [x for x in fogli if x[3]]
    if not fogli:
        flash("Nel file non ho trovato lezioni: non ho cambiato nulla.")
        return redirect(url_for("ripetizioni"))
    try:
        # il nuovo file SOSTITUISCE tutto: lezioni (anche quelle inserite a mano), pagamenti, crediti, note e allievi
        Lezione.query.filter_by(user_id=uid).delete()
        esistenti = {s.nome: s for s in Studente.query.filter_by(user_id=uid)}
        n = nuovi = 0
        nel_file = set()
        for nome, tariffa, nota, lez in fogli:
            s = esistenti.get(nome)
            if not s:
                s = Studente(user_id=uid, nome=nome)
                db.session.add(s)
                nuovi += 1
            # dati dell'allievo presi solo dal file (se ne' tariffa ne' credito sono nel file, ripartono da zero)
            s.tariffa, s.note = tariffa, nota
            s.credito = crediti[nome] if crediti.get(nome) is not None else da_note.get(nome, 0)
            db.session.flush()
            nel_file.add(s.id)
            db.session.add_all(Lezione(user_id=uid, studente_id=s.id, fonte="import", **l) for l in lez)
            n += len(lez)
        rimossi = Studente.query.filter(Studente.user_id == uid, ~Studente.id.in_(nel_file)).delete(synchronize_session=False)  # allievi non piu' nel file
        rec = FileCaricato.query.filter_by(user_id=uid, tipo="ripetizioni").first() or FileCaricato(user_id=uid, tipo="ripetizioni")
        rec.nome, rec.contenuto = (f.filename or "Ripetizioni.xlsx")[:200], dati
        rec.caricato_il = datetime.now(timezone.utc).replace(tzinfo=None)
        db.session.add(rec)
        db.session.commit()
    except Exception:
        db.session.rollback()
        flash("Errore durante l'importazione: i dati precedenti non sono stati modificati.")
        return redirect(url_for("ripetizioni"))
    flash(f"File caricato: i dati precedenti sono stati sostituiti. Ora ci sono {n} lezioni e {len(fogli)} allievi ({nuovi} nuovi, {rimossi} rimossi).")
    return redirect(url_for("ripetizioni"))


@app.route("/modifica-lezione/<int:id>", methods=["GET", "POST"])
@login_required
def rip_modifica(id):
    l = mio(Lezione, id)
    if request.method == "POST":
        f = request.form
        try:
            l.data = datetime.strptime(f["data"], "%Y-%m-%d").date()
            l.ore, l.tariffa = num(f["ore"]), abs(num(f["tariffa"]))
            if l.ore <= 0:
                raise ValueError
            l.ora_inizio, l.ora_fine = f.get("inizio", ""), f.get("fine", "")
            l.materia, l.modalita = f.get("materia", "").strip(), f.get("modalita", "")
            nuova = f.get("pagata") == "1"
            if nuova != l.pagata:
                l.pagata, l.data_pagamento = nuova, (date.today() if nuova else None)
            db.session.commit()
            return redirect(url_for("allievo", id=l.studente_id))
        except (ValueError, KeyError):
            db.session.rollback()
            flash("Controlla i dati inseriti.")
            return redirect(url_for("rip_modifica", id=id))
    return render_template("lezione.html", l=l)


@app.route("/ripetizioni/confronti")
@login_required
def rip_confronti():
    dati = {}  # "AAAA-MM" -> [lezioni, ore, incassato]
    q = (db.session.query(Lezione.data, Lezione.ore, Lezione.tariffa, Lezione.pagata, Studente.tariffa)
         .join(Studente, Lezione.studente_id == Studente.id).filter(Lezione.user_id == session["uid"]))
    for d, ore, tar, pag, st in q:
        r = dati.setdefault(d.strftime("%Y-%m"), [0, 0.0, 0.0])
        r[0] += 1
        r[1] = round(r[1] + ore, 2)
        if pag:
            r[2] = round(r[2] + ore * (tar or st or 0), 2)
    return render_template("confronti.html", dati=dati)


MESI_LUNGHI = list(MESI_IT)


@app.template_filter("ore")
def filtro_ore(x):
    return f"{x:.2f}".rstrip("0").rstrip(".")


def ore_tra(uid, a, b):
    """Ore di lezione con data da a (incluso) a b (escluso)."""
    return db.session.query(db.func.coalesce(db.func.sum(Lezione.ore), 0)).filter(
        Lezione.user_id == uid, Lezione.data >= a, Lezione.data < b).scalar() or 0


def ore_periodi(uid):
    oggi, g = date.today(), timedelta(days=1)
    m0 = oggi.replace(day=1)
    pm0 = (m0 - g).replace(day=1)
    m1 = (m0 + timedelta(days=32)).replace(day=1)
    s0 = oggi - timedelta(days=oggi.weekday())
    ps0 = s0 - timedelta(days=7)
    fine_ps = s0 - g
    return dict(
        mese=ore_tra(uid, m0, oggi + g), mese_prog=ore_tra(uid, oggi + g, m1),
        mese_prec=ore_tra(uid, pm0, m0), mese_prec_pari=ore_tra(uid, pm0, min(m0, pm0 + timedelta(days=oggi.day))),
        sett=ore_tra(uid, s0, oggi + g), sett_prog=ore_tra(uid, oggi + g, s0 + timedelta(days=7)),
        sett_prec=ore_tra(uid, ps0, s0), sett_prec_pari=ore_tra(uid, ps0, ps0 + timedelta(days=oggi.weekday() + 1)),
        nome_mese_prec=f"{MESI_LUNGHI[pm0.month - 1]} {pm0.year}",
        nome_mese_corr=f"{MESI_LUNGHI[m0.month - 1]} {m0.year}",
        periodo_sett_prec=f"dal {ps0.day}/{ps0.month} al {fine_ps.day}/{fine_ps.month}",
        periodo_sett=f"dal {s0.day}/{s0.month} a oggi", n_sett=oggi.isocalendar()[1], n_sett_prec=ps0.isocalendar()[1])


@app.template_filter("ore_it")
def filtro_ore_it(x):
    return filtro_ore(x).replace(".", ",")


@app.context_processor
def _util_pct():
    from markupsafe import Markup
    def pct(cur, prec):
        """Variazione percentuale tra due importi, in piccolo tra parentesi (vuoto se il precedente e' zero)."""
        if not prec:
            return ""
        d = round((cur / prec - 1) * 100)
        return Markup(f'<span class="pc">({"+" if d >= 0 else "−"}{abs(d)}%)</span>')
    return {"pct": pct}


@app.template_filter("eur_it")
def filtro_eur_it(x):
    """Euro all'italiana: senza decimali se intero, altrimenti 2 (1.105 / 1.105,50)"""
    return filtro_it2(x).removesuffix(",00")


@app.template_filter("it2")
def filtro_it2(x):
    """Numero con 2 decimali all'italiana: 1.234,50"""
    return f"{x:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def tariffe_periodi(uid):
    """Ore e incasso maturato per ogni tariffa oraria nei periodi mese/settimana (solo lezioni gia' svolte) e media euro/ora."""
    oggi, g = date.today(), timedelta(days=1)
    m0 = oggi.replace(day=1)
    pm0 = (m0 - g).replace(day=1)
    s0 = oggi - timedelta(days=oggi.weekday())
    ps0 = s0 - timedelta(days=7)
    periodi = {"mese": (m0, oggi + g), "mese_prec": (pm0, m0), "sett": (s0, oggi + g), "sett_prec": (ps0, s0),
               # stessi giorni del periodo precedente (per il colore verde/rosso del fatturato)
               "mese_prec_pari": (pm0, min(m0, pm0 + timedelta(days=oggi.day))), "sett_prec_pari": (ps0, ps0 + timedelta(days=oggi.weekday() + 1))}
    out = {}
    for k, (a, b) in periodi.items():
        righe = (db.session.query(Lezione.ore, Lezione.tariffa, Studente.tariffa)
                 .join(Studente, Studente.id == Lezione.studente_id)
                 .filter(Lezione.user_id == uid, Lezione.data >= a, Lezione.data < b).all())
        per = {}
        for ore, tl, ts in righe:
            t = round(float(tl if tl is not None else (ts or 0)), 2)
            per[t] = per.get(t, 0.0) + (ore or 0)
        voci = [{"t": t, "ore": round(o, 2), "euro": round(o * t, 2)} for t, o in sorted(per.items(), reverse=True) if o]
        ore_t, euro_t = round(sum(v["ore"] for v in voci), 2), round(sum(v["euro"] for v in voci), 2)
        out[k] = {"voci": voci, "ore": ore_t, "euro": euro_t, "media": round(euro_t / ore_t, 2) if ore_t else None}
    return out


@app.route("/ripetizioni/calendario")
@login_required
def rip_calendario():
    """Il vecchio calendario e' stato sostituito dall'Agenda: i vecchi link portano li'."""
    return redirect("/agenda")


def ultimo_file(uid):
    r = (db.session.query(FileCaricato.nome, FileCaricato.caricato_il)
         .filter(FileCaricato.user_id == uid, FileCaricato.tipo == "ripetizioni").first())
    return dict(nome=r.nome, iso=r.caricato_il.isoformat() + "Z") if r else None


@app.route("/ripetizioni/file")
@login_required
def rip_file():
    r = FileCaricato.query.filter_by(user_id=session["uid"], tipo="ripetizioni").first_or_404()
    return send_file(io.BytesIO(r.contenuto), as_attachment=True, download_name=r.nome)


@app.route("/ripetizioni/file/elimina", methods=["POST"])
@login_required
def rip_file_elimina():
    """Toglie il file caricato e i dati che arrivano da lui (lezioni importate e allievi rimasti senza lezioni).
    Le lezioni inserite a mano dal sito non vengono toccate."""
    uid = session["uid"]
    rec = FileCaricato.query.filter_by(user_id=uid, tipo="ripetizioni").first_or_404()
    nome = rec.nome
    ids = {i for (i,) in db.session.query(Lezione.studente_id).filter_by(user_id=uid, fonte="import").distinct()}
    n = Lezione.query.filter_by(user_id=uid, fonte="import").delete()
    vuoti = [i for i in ids if not Lezione.query.filter_by(studente_id=i).count()]
    if vuoti:
        Studente.query.filter(Studente.id.in_(vuoti)).delete(synchronize_session=False)
    db.session.delete(rec)
    db.session.commit()
    flash(f"File «{nome}» eliminato, con le sue {n} lezioni e {len(vuoti)} allievi. Le lezioni inserite a mano non sono state toccate.")
    return redirect(url_for("ripetizioni"))


@app.route("/ripetizioni/esporta-fogli")
@login_required
def rip_esporta_fogli():
    """Excel nel formato originale: un foglio per allievo, testo rosso/verde, tariffa in fondo. Si puo' ricaricare."""
    def ora(v):
        try:
            h, m = map(int, v.split(":"))
            return time(h, m)
        except (ValueError, AttributeError):
            return None
    wb = Workbook()
    wb.remove(wb.active)
    rosso, verde, grassetto = Font(color="FF0000"), Font(color="00B050"), Font(bold=True)
    usati = set()
    for s in Studente.query.filter_by(user_id=session["uid"]).order_by(Studente.id):
        lez = Lezione.query.filter_by(studente_id=s.id).order_by(Lezione.data, Lezione.ora_inizio, Lezione.id).all()
        if not lez:
            continue
        base = re.sub(r"[\[\]:*?/\\]", "", s.nome).strip()[:31] or "Allievo"
        nome, i = base, 2
        while nome.lower() in usati:
            suff = f" {i}"
            nome, i = base[:31 - len(suff)] + suff, i + 1
        usati.add(nome.lower())
        ws = wb.create_sheet(nome)
        ws.append(["Data", "Ora inizio", "Ora fine", "Ore totali", "Materia", "Presenza/DAD", None, None, None, "Tariffa EUR/h", "Note", "Credito EUR"])
        for c in ws[1]:
            c.font = grassetto
        for l in lez:
            ws.append([l.data, ora(l.ora_inizio), ora(l.ora_fine), l.ore, l.materia or None,
                       {"Presenza": "Presenza", "Online": "DAD"}.get(l.modalita), None, None, None, l.tariffa or s.tariffa])
            r = ws.max_row
            for c in range(1, 7):
                ws.cell(r, c).font = verde if l.pagata else rosso
            ws.cell(r, 1).number_format = "DD/MM/YYYY"
            ws.cell(r, 2).number_format = ws.cell(r, 3).number_format = "HH:MM"
        ws["K2"] = s.note or None
        ws["L2"] = round(s.credito or 0, 2)
        ws.append([])
        ws.append(["Importi/ora", s.tariffa])
        for col, w in zip("ABCDEFJKL", (14, 11, 10, 11, 22, 14, 13, 40, 13)):
            ws.column_dimensions[col].width = w
    if not wb.sheetnames:  # nessuna lezione: un Excel senza fogli non si puo' salvare
        flash("Non ci sono ancora lezioni da esportare.")
        return redirect(url_for("ripetizioni"))
    wb.properties.title = "ripetizioni-export"  # riconosce i file scaricati da qui
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Ripetizioni_aggiornato_{date.today()}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------- stipendi (check cedolini) ----------
def _num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _nota_cella(c):
    if not c.comment:
        return ""
    t = re.sub(r"^\s*Marco Cescon:?\s*", "", c.comment.text.strip())
    return " · ".join(p.strip() for p in t.splitlines() if p.strip())


def leggi_stipendi(wb_val, wb_form):
    """Legge il foglio 'Stipendi' (riga 2 = titoli, dati dalla riga 3). Restituisce l'elenco di dizionari."""
    ws = wb_val["Stipendi"] if "Stipendi" in wb_val.sheetnames else wb_val.worksheets[0]
    wf = wb_form[ws.title]
    out = []
    for r in range(3, ws.max_row + 1):
        a = ws.cell(r, 1).value
        d = {"tipo": "mensile"}
        if isinstance(a, (datetime, date)):
            d["anno"], d["mese"] = a.year, a.month
        elif isinstance(a, str) and a.strip():
            m = re.match(r"\s*tredicesima\W*(\d{2,4})", a, re.I)
            if not m:
                continue
            y = int(m.group(1))
            d.update(tipo="tredicesima", anno=y + 2000 if y < 100 else y, mese=12)
        else:
            continue
        for col, att, _t in STIP_CAMPI:
            d[att] = _num(ws[f"{col}{r}"].value)
        if d["competenze"] is None and d["trattenute"] is None:
            continue
        link = ws[f"T{r}"].value
        d["pdf"] = str(link).strip() if link else None
        note = [_nota_cella(wf[f"{c}{r}"]) for c in "IJNRA"]
        d["note"] = " · ".join(n for n in note if n)[:1000] or None
        out.append(d)
    return out


def stip_dict(s):
    d = {"id": s.id, "anno": s.anno, "mese": s.mese, "tipo": s.tipo, "note": s.note or "", "pdf": s.pdf or ""}
    for _c, a, _t in STIP_CAMPI:
        d[a] = getattr(s, a)
    d["netto"] = round((s.competenze or 0) - (s.trattenute or 0), 2)
    d["fe_saldo"] = None if s.fe_prec is None and s.fe_spett is None else round((s.fe_prec or 0) + (s.fe_spett or 0) - (s.fe_god or 0), 2)
    d["par_saldo"] = None if s.par_prec is None and s.par_spett is None else round((s.par_prec or 0) + (s.par_spett or 0) - (s.par_god or 0), 2)
    d["mensa_gg"] = round(s.mensa_tratt / s.gg_mensa, 2) if s.mensa_tratt and s.gg_mensa else None
    return d


def controlli_stipendi(righe):
    """Controlli di coerenza tra mesi consecutivi di ferie e PAR (A. prec. / Spett. / Godute):
    - a gennaio il riporto dell'anno precedente deve essere il saldo di dicembre;
    - le ore spettanti crescono ogni mese del solito accredito (quello piu' frequente);
    - le ore godute non possono diminuire nell'anno.
    Restituisce {id: [avvisi]}."""
    from collections import Counter
    avv = {}
    mens = [r for r in righe if r["tipo"] == "mensile"]
    gruppi = (("ferie", "fe_prec", "fe_spett", "fe_god", "fe_saldo"), ("PAR", "par_prec", "par_spett", "par_god", "par_saldo"))
    acc = {}
    for nome, _p, sp, _g, _s in gruppi:
        c = Counter()
        for prec, r in zip(mens, mens[1:]):
            if r["anno"] == prec["anno"] and r["mese"] == prec["mese"] + 1 and r[sp] is not None and prec[sp] is not None:
                c[round(r[sp] - prec[sp], 2)] += 1
        acc[nome] = c.most_common(1)[0][0] if c else None
    def add(r, t):
        avv.setdefault(r["id"], []).append(t)
    for prec, r in zip(mens, mens[1:]):
        if (r["anno"], r["mese"]) != (prec["anno"] + (prec["mese"] == 12), prec["mese"] % 12 + 1):
            continue  # mese mancante: nessun confronto
        for nome, pr, sp, gd, sl in gruppi:
            if r["mese"] == 1:
                if prec[sl] is not None and r[pr] is not None and abs(prec[sl] - r[pr]) > 0.05:
                    add(r, f"{nome[0].upper() + nome[1:]}: il riporto ({r[pr]:g} h) non coincide con il saldo di dicembre ({prec[sl]:g} h)")
                continue
            if acc[nome] is not None and prec[sp] is not None and r[sp] is not None and abs((r[sp] - prec[sp]) - acc[nome]) > 0.3:
                add(r, f"{nome[0].upper() + nome[1:]}: maturate {r[sp] - prec[sp]:g} h invece delle solite {acc[nome]:g} h")
            if prec[gd] is not None and r[gd] is not None and r[gd] < prec[gd] - 0.05:
                add(r, f"{nome[0].upper() + nome[1:]}: le ore godute sono scese ({prec[gd]:g} → {r[gd]:g} h)")
    costi = [r["mensa_gg"] for r in mens if r["mensa_gg"]]
    if len(costi) >= 6:
        med = sorted(costi)[len(costi) // 2]
        for r in mens:
            if r["mensa_gg"] and abs(r["mensa_gg"] - med) > med * 0.15:
                add(r, f"Mensa: {r['mensa_gg']:.2f} € al giorno (di solito {med:.2f} €)")
    return avv


MESI_LUN = [m.capitalize() for m in MESI_IT]


@app.route("/stipendi")
@login_required
def stipendi():
    uid = session["uid"]
    righe = [stip_dict(s) for s in Stipendio.query.filter_by(user_id=uid).order_by(Stipendio.anno, Stipendio.mese, Stipendio.tipo)]
    avv = controlli_stipendi(righe)
    for r in righe:
        r["avvisi"] = avv.get(r["id"], [])
    f = FileCaricato.query.filter_by(user_id=uid, tipo="stipendi").first()
    ultimo = dict(nome=f.nome, iso=f.caricato_il.replace(tzinfo=timezone.utc).isoformat()) if f else None
    pdfs = {}
    for p in CedolinoPdf.query.filter_by(user_id=uid).with_entities(CedolinoPdf.id, CedolinoPdf.anno, CedolinoPdf.mese, CedolinoPdf.tipo, CedolinoPdf.nome):
        pdfs[f"{p.anno}-{p.mese}-{p.tipo}"] = {"id": p.id, "nome": p.nome, "n": None}
    for r in righe:
        p = pdfs.get(f"{r['anno']}-{r['mese']}-{r['tipo']}")
        if p:
            rec = db.session.get(CedolinoPdf, p["id"])
            p["n"] = sum(1 for x in confronta_pdf(rec, righe) if not x["ok"])
            r["pdf_id"], r["pdf_n"] = p["id"], p["n"]
    elenco = sorted(({"id": p["id"], "nome": p["nome"], "k": k, "n": p["n"]} for k, p in pdfs.items()), key=lambda x: [int(i) if i.isdigit() else i for i in x["k"].split("-")], reverse=True)
    peso = db.session.query(db.func.coalesce(db.func.sum(db.func.length(CedolinoPdf.contenuto)), 0)).filter_by(user_id=uid).scalar() or 0
    return render_template("stipendi.html", righe=righe, ultimo_file=ultimo, mesi=MESI_LUN, pdfs=elenco, peso_pdf=peso,
                           campi=[(a, t) for _c, a, t in STIP_CAMPI])


@app.route("/stipendi/importa", methods=["POST"])
@login_required
def stipendi_importa():
    uid = session["uid"]
    f = request.files.get("file")
    dati = f.read() if f else b""
    try:
        righe = leggi_stipendi(load_workbook(io.BytesIO(dati), data_only=True), load_workbook(io.BytesIO(dati)))
    except Exception:
        flash("Il file non e' un Excel (.xlsx) valido: non ho cambiato nulla.")
        return redirect(url_for("stipendi"))
    if not righe:
        flash("Non ho trovato righe di stipendi nel file (serve il foglio «Stipendi» con le date dei mesi nella colonna A): non ho cambiato nulla.")
        return redirect(url_for("stipendi"))
    Stipendio.query.filter_by(user_id=uid).delete()
    for d in righe:
        db.session.add(Stipendio(user_id=uid, fonte="import", **d))
    FileCaricato.query.filter_by(user_id=uid, tipo="stipendi").delete()
    db.session.add(FileCaricato(user_id=uid, tipo="stipendi", nome=(f.filename or "Stipendi.xlsx")[:200],
                                caricato_il=datetime.now(timezone.utc).replace(tzinfo=None), contenuto=dati))
    db.session.commit()
    flash(f"Importati {len(righe)} cedolini.")
    return redirect(url_for("stipendi"))


@app.route("/stipendi/file")
@login_required
def stipendi_file():
    r = FileCaricato.query.filter_by(user_id=session["uid"], tipo="stipendi").first_or_404()
    return send_file(io.BytesIO(r.contenuto), as_attachment=True, download_name=r.nome)


@app.route("/stipendi/file/elimina", methods=["POST"])
@login_required
def stipendi_file_elimina():
    uid = session["uid"]
    rec = FileCaricato.query.filter_by(user_id=uid, tipo="stipendi").first_or_404()
    n = Stipendio.query.filter_by(user_id=uid, fonte="import").delete()
    nome = rec.nome
    db.session.delete(rec)
    db.session.commit()
    flash(f"File «{nome}» eliminato, con i suoi {n} cedolini. I mesi inseriti a mano dal sito non sono stati toccati.")
    return redirect(url_for("stipendi"))


@app.route("/stipendi/salva", methods=["POST"])
@login_required
def stipendi_salva():
    uid = session["uid"]
    try:
        anno, mese = int(request.form["anno"]), int(request.form["mese"])
        assert 2000 <= anno <= 2100 and 1 <= mese <= 12
    except (KeyError, ValueError, AssertionError):
        flash("Anno o mese non validi.")
        return redirect(url_for("stipendi"))
    tipo = "tredicesima" if request.form.get("tipo") == "tredicesima" else "mensile"
    s = Stipendio.query.filter_by(user_id=uid, anno=anno, mese=mese, tipo=tipo).first()
    nuovo = s is None
    if nuovo:
        s = Stipendio(user_id=uid, anno=anno, mese=mese, tipo=tipo, fonte="sito")
        db.session.add(s)
    for _c, a, _t in STIP_CAMPI:
        v = request.form.get(a, "").strip().replace(",", ".")
        setattr(s, a, _num(v))
    s.note = request.form.get("note", "").strip()[:1000] or None
    s.pdf = request.form.get("pdf", "").strip()[:200] or None
    db.session.commit()
    flash(f"Cedolino {'aggiunto' if nuovo else 'aggiornato'}: {MESI_LUN[mese - 1]} {anno}{' (tredicesima)' if tipo == 'tredicesima' else ''}.")
    return redirect(url_for("stipendi"))


@app.route("/stipendi/<int:id>/elimina", methods=["POST"])
@login_required
def stipendi_elimina(id):
    s = Stipendio.query.filter_by(id=id, user_id=session["uid"]).first_or_404()
    db.session.delete(s)
    db.session.commit()
    flash("Cedolino eliminato.")
    return redirect(url_for("stipendi"))


@app.route("/stipendi/esporta")
@login_required
def stipendi_esporta():
    """Excel nello stesso formato di Check_stipendi (si puo' ricaricare)."""
    righe = Stipendio.query.filter_by(user_id=session["uid"]).order_by(Stipendio.anno, Stipendio.mese, Stipendio.tipo).all()
    nomi_pdf = {(p.anno, p.mese, p.tipo): p.nome for p in CedolinoPdf.query.filter_by(user_id=session["uid"]).with_entities(CedolinoPdf.anno, CedolinoPdf.mese, CedolinoPdf.tipo, CedolinoPdf.nome)}
    if not righe:
        flash("Non ci sono ancora cedolini da esportare.")
        return redirect(url_for("stipendi"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Stipendi"
    ws["L1"], ws["P1"] = "FERIE", "PAR"
    tit = {"K": "NETTO", "O": "Saldo", "S": "Saldo", "T": "LINK PDF", "Y": "Costo mensa giornaliero medio"}
    tit.update({c: t for c, _a, t in STIP_CAMPI})
    for c, t in tit.items():
        ws[f"{c}2"] = t
    for c in ws[2]:
        c.font = Font(bold=True)
    for i, s in enumerate(righe):
        r = i + 3
        if s.tipo == "tredicesima":
            ws[f"A{r}"] = f"tredicesima-{str(s.anno)[2:]}"
        else:
            ws[f"A{r}"] = date(s.anno, s.mese, 1)
            ws[f"A{r}"].number_format = "mmmm yyyy"
        for c, a, _t in STIP_CAMPI:
            ws[f"{c}{r}"] = getattr(s, a)
        ws[f"K{r}"] = f"=I{r}-J{r}"
        ws[f"O{r}"] = f"=L{r}+M{r}-N{r}"
        ws[f"S{r}"] = f"=P{r}+Q{r}-R{r}"
        ws[f"Y{r}"] = f'=IF(AND(C{r}>0,W{r}<>""),W{r}/C{r},"")'
        ws[f"T{r}"] = s.pdf or nomi_pdf.get((s.anno, s.mese, s.tipo))
        if s.note:
            ws[f"I{r}"].comment = Comment(s.note, "Le mie spese")
    ws.column_dimensions["A"].width = 16
    ws.freeze_panes = "B3"
    wb.properties.title = "stipendi-export"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"Check_stipendi_aggiornato_{date.today()}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------- cedolini PDF ----------
PDF_CAMPI = [(a, t) for _c, a, t in STIP_CAMPI if a not in ("mensa_sup",)]


def confronta_pdf(rec, righe=None):
    """Confronta i valori letti dal PDF con la riga di Check_stipendi (stesso mese). Restituisce un elenco di righe."""
    v = json.loads(rec.dati)["v"]
    if righe is None:
        righe = [stip_dict(s) for s in Stipendio.query.filter_by(user_id=rec.user_id).order_by(Stipendio.anno, Stipendio.mese)]
    riga = next((r for r in righe if (r["anno"], r["mese"], r["tipo"]) == (rec.anno, rec.mese, rec.tipo)), None)
    out = []
    if not riga:
        return out
    campi = [(a, t) for a, t in PDF_CAMPI]
    campi.append(("netto", "Netto da pagare"))
    for a, t in campi:
        p, x = v.get(a), riga.get(a)
        if a in ("fe_god", "par_god"):
            p, x = p or 0, x or 0
        if p is None and x is None:
            continue
        if p is None:
            continue  # il PDF non riporta quel dato: non si puo' confrontare
        tol = 0.001 if a == "unitario" else 0.011
        out.append({"campo": a, "titolo": t, "pdf": p, "riga": x, "ok": x is not None and abs(p - x) <= tol})
    # ore scalate nel mese, calcolate dai saldi: saldo del mese prima + ore maturate nel mese - saldo di questo mese.
    # (non basta guardare la colonna "Godute": a volte le ore vengono tolte dal riporto dell'anno precedente)
    prec = next((r for r in righe if r["tipo"] == "mensile" and (r["anno"], r["mese"]) == ((rec.anno, rec.mese - 1) if rec.mese > 1 else (rec.anno - 1, 12))), None)
    if rec.tipo == "mensile" and prec:
        for pre, nome, chiave in (("fe", "Ferie", "_fe_desc"), ("par", "PAR", "_par_desc")):
            ps, sp, sl_prec = v.get(f"{pre}_prec"), v.get(f"{pre}_spett"), prec.get(f"{pre}_saldo")
            if None in (ps, sp, sl_prec):
                continue
            sl = ps + sp - (v.get(f"{pre}_god") or 0)
            maturate = sp - ((prec.get(f"{pre}_spett") or 0) if rec.mese > 1 else 0)
            scalate = round(sl_prec + maturate - sl, 2)
            desc = v.get(chiave) or 0.0
            out.append({"campo": f"_{pre}_mese", "titolo": f"{nome} del mese: ore in descrizione / ore scalate dal saldo", "pdf": desc, "riga": scalate, "ok": abs(desc - scalate) <= 0.011,
                        "nota": f"Saldo {MESI_LUN[prec['mese'] - 1].lower()} {prec.get(f'{pre}_saldo'):g} h + maturate nel mese {maturate:g} h − saldo di questo mese {sl:g} h = {scalate:g} h scalate"})
    return out


def _applica_valori(uid, v):
    s = Stipendio.query.filter_by(user_id=uid, anno=v["anno"], mese=v["mese"], tipo=v["tipo"]).first()
    nuovo = s is None
    if nuovo:
        s = Stipendio(user_id=uid, anno=v["anno"], mese=v["mese"], tipo=v["tipo"], fonte="pdf")
        db.session.add(s)
    for a, _t in PDF_CAMPI:
        if v.get(a) is not None or a in ("fe_god", "par_god"):
            setattr(s, a, v.get(a))
    return nuovo


@app.route("/stipendi/pdf", methods=["POST"])
@login_required
def stipendi_pdf():
    from cedolino import leggi_cedolino
    uid = session["uid"]
    msg = []
    for f in request.files.getlist("file"):
        dati = f.read()
        if not dati:
            continue
        try:
            v, voci, ctrl = leggi_cedolino(dati)
        except ValueError as e:
            msg.append(f"{f.filename}: {e}.")
            continue
        rec = CedolinoPdf.query.filter_by(user_id=uid, anno=v["anno"], mese=v["mese"], tipo=v["tipo"]).first()
        if not rec:
            rec = CedolinoPdf(user_id=uid, anno=v["anno"], mese=v["mese"], tipo=v["tipo"])
            db.session.add(rec)
        rec.nome = (f.filename or "cedolino.pdf")[:200]
        rec.caricato_il = datetime.now(timezone.utc).replace(tzinfo=None)
        rec.contenuto = dati
        rec.dati = json.dumps({"v": v, "voci": voci, "c": ctrl})
        db.session.flush()
        creato = False
        if not Stipendio.query.filter_by(user_id=uid, anno=v["anno"], mese=v["mese"], tipo=v["tipo"]).first():
            creato = _applica_valori(uid, v)
        db.session.flush()
        n = sum(1 for x in confronta_pdf(rec) if not x["ok"]) + sum(1 for x in ctrl if not x["ok"])
        msg.append(f"{MESI_LUN[v['mese'] - 1]} {v['anno']}: {'riga creata dal PDF, ' if creato else ''}" + ("tutto coincide." if not n else f"{n} da controllare."))
    db.session.commit()
    flash(" ".join(msg) or "Nessun file ricevuto.")
    return redirect(url_for("stipendi"))


@app.route("/stipendi/pdf/rileggi", methods=["POST"])
@login_required
def stipendi_pdf_rileggi():
    """Rilegge tutti i PDF salvati con il lettore attuale (utile dopo un miglioramento del lettore)."""
    from cedolino import leggi_cedolino
    n = err = 0
    for rec in CedolinoPdf.query.filter_by(user_id=session["uid"]):
        try:
            v, voci, ctrl = leggi_cedolino(rec.contenuto)
        except ValueError:
            err += 1
            continue
        rec.dati = json.dumps({"v": v, "voci": voci, "c": ctrl})
        n += 1
    db.session.commit()
    flash(f"Riletti {n} PDF" + (f"; {err} non leggibili." if err else "."))
    return redirect(url_for("stipendi"))


@app.route("/stipendi/pdf/<int:id>")
@login_required
def stipendi_pdf_dettaglio(id):
    rec = CedolinoPdf.query.filter_by(id=id, user_id=session["uid"]).first_or_404()
    d = json.loads(rec.dati)
    righe = [stip_dict(s) for s in Stipendio.query.filter_by(user_id=rec.user_id)]
    return render_template("cedolino.html", rec=rec, conf=confronta_pdf(rec, righe), voci=d["voci"], ctrl=d["c"],
                           ha_riga=any((r["anno"], r["mese"], r["tipo"]) == (rec.anno, rec.mese, rec.tipo) for r in righe),
                           mese=MESI_LUN[rec.mese - 1], iso=rec.caricato_il.replace(tzinfo=timezone.utc).isoformat())


@app.route("/stipendi/pdf/<int:id>/file")
@login_required
def stipendi_pdf_file(id):
    rec = CedolinoPdf.query.filter_by(id=id, user_id=session["uid"]).first_or_404()
    return send_file(io.BytesIO(rec.contenuto), as_attachment=True, download_name=rec.nome, mimetype="application/pdf")


@app.route("/stipendi/pdf/<int:id>/elimina", methods=["POST"])
@login_required
def stipendi_pdf_elimina(id):
    rec = CedolinoPdf.query.filter_by(id=id, user_id=session["uid"]).first_or_404()
    db.session.delete(rec)
    db.session.commit()
    flash("PDF eliminato (la riga del cedolino resta).")
    return redirect(url_for("stipendi"))


@app.route("/stipendi/pdf/<int:id>/applica", methods=["POST"])
@login_required
def stipendi_pdf_applica(id):
    rec = CedolinoPdf.query.filter_by(id=id, user_id=session["uid"]).first_or_404()
    _applica_valori(session["uid"], json.loads(rec.dati)["v"])
    db.session.commit()
    flash("Riga aggiornata con i valori del PDF.")
    return redirect(url_for("stipendi_pdf_dettaglio", id=id))



# ---------- Stipendi: ore giorno per giorno ----------
ORE_STORICO = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ore_storico_2026.json")
GIORNI_IT = ["lun", "mar", "mer", "gio", "ven", "sab", "dom"]


def ore_dict(o):
    d = {"data": o.data.isoformat(), "pranzo": o.pranzo or "", "pranzo_nota": o.pranzo_nota or "", "nota": o.nota or ""}
    for a in ORE_CAMPI:
        d[a] = getattr(o, a)
    # "Extra / spese varie" e' una sola colonna: include il vecchio importo del pranzo
    ex = round((o.spese or 0) + (o.pranzo_importo or 0), 2)
    d["spese"], d["pranzo_importo"] = (ex or None), None
    return d


def mese_param(txt):
    m = re.fullmatch(r"(\d{4})-(\d{1,2})", txt or "")
    if m and 2000 <= int(m.group(1)) <= 2100 and 1 <= int(m.group(2)) <= 12:
        return int(m.group(1)), int(m.group(2))
    oggi = date.today()
    return oggi.year, oggi.month


def giorni_feriali(anno, mese):
    import calendar
    n = calendar.monthrange(anno, mese)[1]
    return sum(1 for g in range(1, n + 1) if date(anno, mese, g).weekday() < 5)


def riepilogo_ore(uid, anno, mese):
    """Totali del mese e confronto con il cedolino (PDF se c'e', altrimenti la riga di Check_stipendi)."""
    import calendar
    n = calendar.monthrange(anno, mese)[1]
    righe = OreGiorno.query.filter(OreGiorno.user_id == uid, OreGiorno.data >= date(anno, mese, 1), OreGiorno.data <= date(anno, mese, n)).all()
    t = {a: round(sum(getattr(r, a) or 0 for r in righe), 2) for a in ORE_CAMPI}
    t["mensa"] = sum(1 for r in righe if r.pranzo == "mensa")
    t["fuori"] = sum(1 for r in righe if r.pranzo == "fuori")
    ex = lambda r: (r.pranzo_importo or 0) + (r.spese or 0)  # extra / spese varie del giorno
    t["extra"] = round(sum(ex(r) for r in righe), 2)
    t["fuori_eur"] = round(sum(ex(r) for r in righe if r.pranzo == "fuori"), 2)
    t["giorni_reg"] = len(righe)
    t["feriali"] = giorni_feriali(anno, mese)
    t["giorni"] = round(t["feriali"] - t["ferie"] / 8 - t["fest_non"] / 8, 2)
    t["sup"] = round(sum(ex(r) for r in righe if r.pranzo == "mensa"), 2)
    conf, fonte = [], None
    srow = Stipendio.query.filter_by(user_id=uid, anno=anno, mese=mese, tipo="mensile").first()
    rec = CedolinoPdf.query.filter_by(user_id=uid, anno=anno, mese=mese, tipo="mensile").first()
    if rec and rec.dati:
        v = json.loads(rec.dati)["v"]
        ev, leg = v.get("_ev") or {}, v.get("_leg") or {}
        def da(*cands):
            for c in cands:
                if c is not None:
                    return c
            return None
        pdf = {
            "lavoro": da(leg.get("OR"), ev.get("OR")),
            "sw": da(v.get("_sw_ore"), leg.get("WS"), ev.get("WS")),
            "ferie": da(v.get("_fe_desc"), leg.get("FR"), ev.get("FR")),
            "permessi": da(v.get("_par_desc"), leg.get("PA"), ev.get("PA")),
            "fest_non": da(leg.get("FS"), ev.get("FS")),
            "trasferta": da(leg.get("SE"), ev.get("SE")),
            "giorni": v.get("giorni"), "mensa": v.get("gg_mensa")}
        fonte = f"cedolino PDF di {MESI_LUN[mese - 1].lower()} {anno}"
        ore_cat = True
    else:
        s = Stipendio.query.filter_by(user_id=uid, anno=anno, mese=mese, tipo="mensile").first()
        pdf = {"giorni": s.giorni if s else None, "mensa": s.gg_mensa if s else None}
        fonte = "riga di Check_stipendi (manca il PDF del cedolino)" if s and (s.giorni is not None or s.gg_mensa is not None) else None
        ore_cat = False
    for k, nome, um in (("lavoro", "Ore di lavoro in sede", "h"), ("sw", "Smart working", "h"), ("ferie", "Ferie", "h"), ("permessi", "Permessi", "h"),
                        ("fest_non", "Festività non godute", "h"), ("trasferta", "Trasferta", "h"), ("giorni", "Giorni lavorati", "gg"), ("mensa", "Giorni di mensa", "gg")):
        p = pdf.get(k)
        sito = t["giorni"] if k == "giorni" else t["mensa"] if k == "mensa" else t[k]
        if k in ("lavoro", "sw", "ferie", "permessi", "fest_non", "trasferta") and not ore_cat:
            continue
        if p is None and k in ("giorni", "mensa"):
            continue
        p = p or 0
        if t["giorni_reg"] == 0:
            continue
        conf.append({"k": k, "nome": nome, "um": um, "sito": sito, "pdf": p, "ok": abs(sito - p) < 0.011})
    if srow is not None and srow.mensa_sup is not None and t["giorni_reg"]:
        conf.append({"k": "sup", "nome": "Supplemento mensa", "um": "€", "sito": t["sup"], "pdf": srow.mensa_sup, "ok": abs(t["sup"] - srow.mensa_sup) < 0.011})
    return {"tot": t, "conf": conf, "fonte": fonte, "pdf": bool(rec)}


@app.route("/stipendi/ore")
@login_required
def ore_pagina():
    uid = session["uid"]
    anno, mese = mese_param(request.args.get("mese"))
    n = OreGiorno.query.filter_by(user_id=uid).count()
    return render_template("ore.html", anno=anno, mese=mese, vuoto=(n == 0), storico=os.path.exists(ORE_STORICO), mesi=MESI_LUN)


@app.route("/stipendi/ore/dati")
@login_required
def ore_dati():
    import calendar
    uid = session["uid"]
    anno, mese = mese_param(request.args.get("mese"))
    n = calendar.monthrange(anno, mese)[1]
    righe = {r.data.day: ore_dict(r) for r in OreGiorno.query.filter(OreGiorno.user_id == uid, OreGiorno.data >= date(anno, mese, 1), OreGiorno.data <= date(anno, mese, n))}
    giorni = [{"g": g, "data": date(anno, mese, g).isoformat(), "wd": GIORNI_IT[date(anno, mese, g).weekday()], "r": righe.get(g)} for g in range(1, n + 1)]
    return jsonify(giorni=giorni, **riepilogo_ore(uid, anno, mese))


@app.route("/stipendi/ore/salva", methods=["POST"])
@login_required
def ore_salva():
    uid = session["uid"]
    try:
        d = date.fromisoformat(request.form.get("data", ""))
    except ValueError:
        abort(400)
    o = OreGiorno.query.filter_by(user_id=uid, data=d).first()
    vals = {}
    for a in ORE_CAMPI:
        x = _num(request.form.get(a, "").strip().replace(",", "."))
        vals[a] = x if x is not None and 0 <= x <= 100000 else None
    pr = request.form.get("pranzo", "")
    pr = pr if pr in ("mensa", "fuori") else None
    pnota, nota = request.form.get("pranzo_nota", "").strip()[:200], request.form.get("nota", "").strip()[:300]
    vuoto = not any(v for v in vals.values()) and not pr and not nota and not pnota
    if vuoto:
        if o:
            db.session.delete(o)
    else:
        if not o:
            o = OreGiorno(user_id=uid, data=d)
            db.session.add(o)
        for a, x in vals.items():
            setattr(o, a, x)
        o.pranzo, o.pranzo_nota, o.nota = pr, pnota or None, nota or None
    db.session.commit()
    return jsonify(ok=True, **riepilogo_ore(uid, d.year, d.month))


@app.route("/stipendi/ore/storico", methods=["POST"])
@login_required
def ore_storico():
    """Carica una volta lo storico 2026 (dal vecchio Excel delle ore), senza toccare i giorni gia' presenti."""
    uid = session["uid"]
    if not os.path.exists(ORE_STORICO):
        abort(404)
    esist = {r.data for r in OreGiorno.query.filter_by(user_id=uid).with_entities(OreGiorno.data)}
    n = 0
    for r in json.load(open(ORE_STORICO, encoding="utf-8")):
        d = date.fromisoformat(r["d"])
        if d in esist:
            continue
        o = OreGiorno(user_id=uid, data=d, pranzo=r.get("pranzo"), pranzo_nota=r.get("pranzo_nota"), nota=r.get("nota"))
        for a in ORE_CAMPI:
            setattr(o, a, r.get(a))
        db.session.add(o)
        n += 1
    db.session.commit()
    flash(f"Storico caricato: {n} giorni.")
    return redirect(url_for("ore_pagina", mese="2026-10"))


@app.route("/stipendi/ore/esporta")
@login_required
def ore_esporta():
    """Excel con un foglio per mese (come il vecchio 'Riepilogo ore'), il piu' recente per primo."""
    import calendar
    from openpyxl.styles import Alignment, PatternFill
    uid = session["uid"]
    righe = OreGiorno.query.filter_by(user_id=uid).order_by(OreGiorno.data).all()
    if not righe:
        flash("Non ci sono ancora ore da esportare.")
        return redirect(url_for("ore_pagina"))
    per_mese = {}
    for r in righe:
        per_mese.setdefault((r.data.year, r.data.month), {})[r.data.day] = r
    wb = Workbook()
    wb.remove(wb.active)
    intest = ["Data", "Ore lavoro", "Smartworking", "Ferie", "Permessi", "Festività non godute", "Festività godute", "Trasferta", "Pranzo", "Caffè", "Extra / spese varie", "Note"]
    grigio = PatternFill("solid", fgColor="EEF1F6")
    for (a, m) in sorted(per_mese, reverse=True):
        ws = wb.create_sheet(f"{MESI_LUN[m - 1]} {a}")
        ws.append(intest)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(wrap_text=True, vertical="center")
        n = _cal.monthrange(a, m)[1]
        for g in range(1, n + 1):
            d = date(a, m, g)
            o = per_mese[(a, m)].get(g)
            if o:
                if o.pranzo == "mensa":
                    pr = "standard"
                elif o.pranzo == "fuori":
                    pr = o.pranzo_nota or "fuori"
                else:
                    pr = None
                ws.append([d, o.lavoro, o.sw, o.ferie, o.permessi, o.fest_non, o.fest_god, o.trasferta, pr, o.caffe, round((o.spese or 0) + (o.pranzo_importo or 0), 2) or None, o.nota])
            else:
                ws.append([d])
            ws.cell(ws.max_row, 1).number_format = "ddd dd/mm/yyyy"
            if d.weekday() >= 5:
                for c in ws[ws.max_row]:
                    c.fill = grigio
        tot = n + 2
        ws.cell(tot, 1, "TOTALE").font = Font(bold=True)
        for col in "BCDEFGHJK":
            ws[f"{col}{tot}"] = f"=SUM({col}2:{col}{n + 1})"
            ws[f"{col}{tot}"].font = Font(bold=True)
        ws[f"I{tot}"] = f'=COUNTIF(I2:I{n + 1},"standard")+COUNT(I2:I{n + 1})'
        ws[f"I{tot}"].font = Font(bold=True)
        ws.cell(tot + 1, 9, "↑ giorni di mensa").font = Font(italic=True, color="666666")
        ws.column_dimensions["A"].width = 17
        for col in "BCDEFGHIJK":
            ws.column_dimensions[col].width = 14
        ws.column_dimensions["L"].width = 30
        ws.freeze_panes = "B2"
    wb.properties.title = "ore-export"
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return send_file(bio, as_attachment=True, download_name=f"Riepilogo_ore_{date.today().isoformat()}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")



# ======================= AGENDA E SCADENZE =======================
import calendar as _cal

COLORI_IMPEGNO = ["#2563eb", "#0891b2", "#16a34a", "#eab308", "#f97316", "#dc2626", "#ec4899", "#64748b"]
SUGGERIMENTI_SCADENZE = ["Assicurazione auto", "Bollo auto", "Revisione auto", "Assicurazione casa", "Abbonamento",
                         "Tari", "Canone Rai", "Imu", "Tessera / abbonamento", "Altro"]


def aggiungi_mesi(d, n):
    m = d.month - 1 + n
    y, m = d.year + m // 12, m % 12 + 1
    return date(y, m, min(d.day, _cal.monthrange(y, m)[1]))


def _ora(v):
    v = (v or "").strip()
    m = re.fullmatch(r"([01]?\d|2[0-3]):([0-5]\d)", v)
    return f"{int(m.group(1)):02d}:{m.group(2)}" if m else ""


def _data(v):
    try:
        return date.fromisoformat((v or "").strip())
    except ValueError:
        return None


def voci_scadenze(uid):
    """Raggruppa i pagamenti per voce: per ognuna l'ultimo pagamento (quello che conta) e lo storico."""
    gruppi = {}
    for s in Scadenza.query.filter_by(user_id=uid).all():
        if s.tipo == "nota":
            continue
        gruppi.setdefault(s.voce.strip().lower(), []).append(s)
    oggi, voci = date.today(), []
    for lista in gruppi.values():
        lista.sort(key=lambda s: (s.scadenza or date.min, s.data_pagamento or date.min, s.id))
        ult = lista[-1]
        giorni = (ult.scadenza - oggi).days if ult.scadenza else None
        pre = max(30, ult.preavviso or 30)
        if ult.chiusa or giorni is None:
            stato = "chiusa" if ult.chiusa else "nd"
        elif giorni < 0:
            stato = "scaduta"
        elif giorni <= 7:
            stato = "urgente"
        elif giorni <= pre:
            stato = "presto"
        else:
            stato = "ok"
        perc = None
        if ult.data_pagamento and ult.scadenza and ult.scadenza > ult.data_pagamento:
            perc = round(min(1, max(0, (oggi - ult.data_pagamento).days / (ult.scadenza - ult.data_pagamento).days)) * 100)
        voci.append(dict(s=ult, storico=list(reversed(lista)), giorni=giorni, stato=stato, perc=perc))
    ordine = {"scaduta": 0, "urgente": 1, "presto": 2, "ok": 3, "nd": 4, "chiusa": 5}
    voci.sort(key=lambda v: (ordine[v["stato"]], v["s"].scadenza or date.max, v["s"].voce.lower()))
    return voci


def alert_scadenze(uid):
    return [v for v in voci_scadenze(uid) if v["stato"] in ("scaduta", "urgente", "presto")]


def _it(d):
    return f"{d.day} {MESI_LUNGHI[d.month - 1][:3]} {d.year}"


def _icona(voce):
    v = (voce or "").lower()
    for chiavi, ico in ((("assicur",), "🛡️"), (("bollo",), "🧾"), (("caldaia", "camin", "gas", "luce", "energia", "riscald"), "🔥"),
                        (("auto", "revision", "moto", "gomme", "tagliando"), "🚗"), (("casa", "imu", "tari", "affitto", "condomin", "mutuo"), "🏠"),
                        (("abbon", "palestra", "netflix", "spotify", "tessera"), "🎟️"), (("tasse", "irpef", "f24", "canone", "rai"), "🏛️"),
                        (("medic", "dent", "salute", "visita"), "🩺")):
        if any(k in v for k in chiavi):
            return ico
    return "📌"


app.jinja_env.filters["icona"] = _icona
app.jinja_env.filters["data_it"] = lambda d: _it(d) if d else "—"
app.jinja_env.filters["validita"] = lambda m: ("" if not m else f"{m // 12} anno" if m == 12 else f"{m // 12} anni" if m % 12 == 0 else f"{m} mesi")


def _colonne(g):
    """Eventi che si sovrappongono: ognuno la sua colonna, affiancati."""
    fine_col, gruppo, fine_gruppo = [], [], 0
    for e in sorted(g["ev"], key=lambda x: (x["start"], x["end"])):
        vis = max(e["end"], e["start"] + 30)
        if gruppo and e["start"] >= fine_gruppo:
            for x in gruppo:
                x["ncols"] = len(fine_col)
            fine_col, gruppo, fine_gruppo = [], [], 0
        for i, fc in enumerate(fine_col):
            if fc <= e["start"]:
                e["col"], fine_col[i] = i, vis
                break
        else:
            e["col"] = len(fine_col)
            fine_col.append(vis)
        gruppo.append(e)
        fine_gruppo = max(fine_gruppo, vis)
    for x in gruppo:
        x["ncols"] = len(fine_col)


def eventi_periodo(uid, da, a):
    """Tutti gli eventi tra due date (comprese): lezioni, impegni, scadenze. Dict data -> lista."""
    out = {}
    def put(d, e):
        out.setdefault(d, []).append(e)
    q = (db.session.query(Lezione, Studente).join(Studente, Lezione.studente_id == Studente.id)
         .filter(Lezione.user_id == uid, Lezione.data >= da, Lezione.data <= a))
    for l, s in q:
        e = dict(kind="lez", titolo=s.nome, ini=l.ora_inizio or "", fine=l.ora_fine or "", href=f"/ripetizioni/{s.id}",
                 cls="ok" if l.pagata else "no", sub=" ".join(x for x in (l.materia, l.modalita) if x), ore=l.ore)
        put(l.data, e)
    for i in Impegno.query.filter(Impegno.user_id == uid, Impegno.data >= da, Impegno.data <= a):
        e = dict(kind="imp", titolo=i.titolo, ini=i.ora_inizio or "", fine=i.ora_fine or "", href="#", cls="imp", sub=i.luogo or "",
                 raw=dict(id=i.id, data=i.data.isoformat(), ini=i.ora_inizio or "", fine=i.ora_fine or "", titolo=i.titolo,
                          luogo=i.luogo or "", nota=i.nota or "", colore=i.colore or ""),
                 colore=i.colore or "")
        put(i.data, e)
    for v in voci_scadenze(uid):
        s = v["s"]
        if s.scadenza and da <= s.scadenza <= a and not s.chiusa:
            put(s.scadenza, dict(kind="scad", titolo=f"SCADENZA {s.voce}", ini="", fine="", href="/scadenze", cls="scad",
                                 sub=(f"{s.importo:g} €".replace(".", ",") if s.importo else "")))
    for dc in Documento.query.filter(Documento.user_id == uid, Documento.scadenza >= da, Documento.scadenza <= a):
        put(dc.scadenza, dict(kind="scad", titolo=f"SCADENZA {nome_doc(dc)}", ini="", fine="", href="/documenti", cls="scad",
                              sub=dc.intestatario or ""))
    for d, lst in out.items():
        for e in lst:
            try:
                h, m = map(int, e["ini"].split(":"))
                e["start"] = h * 60 + m
                e["end"] = e["start"] + int((e.get("ore") or 1) * 60)
                if e["fine"]:
                    h2, m2 = map(int, e["fine"].split(":"))
                    if h2 * 60 + m2 > e["start"]:
                        e["end"] = h2 * 60 + m2
            except ValueError:
                e["start"] = None
        lst.sort(key=lambda e: (e["start"] is None, e["start"] or 0, e["titolo"]))
    return out


CARTE = ("banc", "cc", "deb")
TIPI_DOC = {
    "ci": "Carta d'identità", "pass": "Passaporto", "pat": "Patente", "nau": "Patente nautica", "ts": "Tessera sanitaria", "cf": "Codice fiscale", "banc": "Bancomat",
    "cc": "Carta di credito", "deb": "Carta di debito", "conto": "Conto corrente", "altro": "Altro documento",
}
COLORI_DOC = ["#2563eb", "#0891b2", "#16a34a", "#eab308", "#f97316", "#dc2626", "#ec4899", "#64748b", "#0f172a"]


def _slug(t):
    import unicodedata
    t = unicodedata.normalize("NFKD", t or "").encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", t).strip("-")


def immagini_doc():
    """File in static/documenti/: nome = <tipo>.ext, <tipo>-<banca>.ext oppure <banca>.ext (png, jpg, jpeg, webp, svg)."""
    base = os.path.join(app.static_folder, "documenti")
    out = {}
    if os.path.isdir(base):
        for f in sorted(os.listdir(base)):
            stem, _, ext = f.rpartition(".")
            if ext.lower() in ("png", "jpg", "jpeg", "webp", "svg") and stem:
                out[_slug(stem)] = f"/static/documenti/{f}?v={int(os.path.getmtime(os.path.join(base, f)))}"
    return out


def immagine_doc(d, imgs):
    """Immagine piu' specifica per un documento: <tipo>-<banca> > <banca> > <tipo>. La banca si riconosce dentro il nome indicato."""
    banca = _slug(d.ente)
    best = None
    for stem, url in imgs.items():
        nome, tipo_f = stem, None
        for t in TIPI_DOC:
            if stem.startswith(t + "-"):
                tipo_f, nome = t, stem[len(t) + 1:]
                break
        if stem in TIPI_DOC or not banca or nome not in banca and nome.replace("-", "") not in banca.replace("-", ""):
            continue
        if tipo_f and tipo_f != d.tipo:
            continue
        punti = len(nome) + (1000 if tipo_f else 0)
        if not best or punti > best[0]:
            best = (punti, url)
    return best[1] if best else imgs.get(d.tipo)


def nome_doc(d):
    """Didascalia: per carte e conti 'BANCA - Tipo' (es. FINECO - Carta di debito)."""
    if d.tipo in ("banc", "cc", "deb", "conto") and (d.ente or "").strip():
        return f"{d.ente.strip().upper()} - {TIPI_DOC[d.tipo]}"
    return d.titolo


_COMUNI = None


def _norm(t):
    import unicodedata
    t = unicodedata.normalize("NFKD", t or "").encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", t).strip()


@app.route("/api/comuni")
@login_required
def api_comuni():
    """Ricerca di comuni italiani o stati esteri (codice catastale) per il calcolo del codice fiscale."""
    global _COMUNI
    if _COMUNI is None:
        with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "comuni.json"), encoding="utf-8") as f:
            _COMUNI = [(_norm(r[0]), r) for r in json.load(f)]
    q = _norm(request.args.get("q", ""))
    if len(q) < 2:
        return jsonify([])
    ini, dentro = [], []
    for n, r in _COMUNI:
        if n.startswith(q):
            ini.append(r)
        elif q in n:
            dentro.append(r)
    ris = sorted(ini, key=lambda r: (not r[3], len(r[0]))) + sorted(dentro, key=lambda r: (not r[3], len(r[0])))
    return jsonify([dict(n=f"{r[0]} ({'estero' if r[2] == 'EE' else r[2]}){'' if r[3] else ' - soppresso'}", c=r[1]) for r in ris[:15]])


def _stato_doc(d, oggi=None):
    """Stato di un documento: giorni alla scadenza e classe (scaduto, urgente <=60 g, presto <=180 g, ok, nd)."""
    oggi = oggi or date.today()
    if not d.scadenza:
        return dict(giorni=None, stato="nd", perc=None)
    g = (d.scadenza - oggi).days
    # barra di validita': dal rilascio (se noto, altrimenti una durata tipica) fino alla scadenza
    anni = {"ci": 10, "pass": 10, "pat": 10, "nau": 10, "ts": 6, "banc": 5, "cc": 5, "deb": 5}.get(d.tipo, 10)
    inizio = d.rilascio if d.rilascio and d.rilascio < d.scadenza else d.scadenza - timedelta(days=round(365.25 * anni))
    tot = (d.scadenza - inizio).days or 1
    perc = max(0, min(100, round((oggi - inizio).days / tot * 100)))
    return dict(giorni=g, stato="scaduta" if g < 0 else "urgente" if g <= 60 else "presto" if g <= 180 else "ok", perc=perc)


def documenti_in_scadenza(uid):
    oggi = date.today()
    return [d for d in Documento.query.filter(Documento.user_id == uid, Documento.scadenza.isnot(None),
                                              Documento.scadenza <= oggi + timedelta(days=60)).order_by(Documento.scadenza)]


def _data_doc(v):
    """Data da un campo date (AAAA-MM-GG) o month (AAAA-MM: vale l'ultimo giorno del mese, come sulle carte)."""
    v = (v or "").strip()
    if len(v) == 7:
        try:
            a, m = int(v[:4]), int(v[5:7])
            return date(a, m, _cal.monthrange(a, m)[1])
        except ValueError:
            return None
    return _data(v)


@app.route("/documenti")
@login_required
def documenti():
    uid = session["uid"]
    docs = Documento.query.filter_by(user_id=uid).order_by(Documento.ordine.is_(None), Documento.ordine, Documento.id).all()
    imgs = immagini_doc()
    lista = []
    per_id = {d.id: d for d in docs}
    carte_di = {}
    for d in docs:
        if d.conto_id in per_id and d.tipo in CARTE:
            carte_di.setdefault(d.conto_id, []).append(nome_doc(d))
    for d in docs:
        st = _stato_doc(d)
        c = per_id.get(d.conto_id) if d.tipo in CARTE else None
        propria = f"/documenti/{d.id}/immagine?v={hashlib.md5(d.immagine.encode()).hexdigest()[:8]}" if d.immagine else None
        lista.append(dict(d=d, nome=nome_doc(d), img=propria or immagine_doc(d, imgs), **st, raw=dict(has_img=bool(d.immagine), 
            id=d.id, tipo=d.tipo, titolo=nome_doc(d), intestatario=d.intestatario or "", nome_p=d.nome_p or "", cognome_p=d.cognome_p or "", numero=d.numero or "", ente=d.ente or "",
            rilascio=d.rilascio.isoformat() if d.rilascio else "", scadenza=d.scadenza.isoformat() if d.scadenza else "",
            iban=d.iban or "", nota=d.nota or "", colore=d.colore or "", giorni=st["giorni"], stato=st["stato"], perc=st.get("perc"), conto_id=c.id if c else "",
            conto=dict(titolo=nome_doc(c), intestatario=c.intestatario or "", iban=c.iban or "", ente=c.ente or "", numero=c.numero or "") if c else None,
            carte=carte_di.get(d.id, []))))
    gruppi = [("Documenti", "doc", ("ci", "pass", "pat", "nau", "ts", "cf", "altro")), ("Carte e dati bancari", "ban", ("banc", "cc", "deb", "conto"))]
    return render_template("documenti.html", gruppi=[(t, k, [x for x in lista if x["d"].tipo in tp]) for t, k, tp in gruppi],
                           tipi=TIPI_DOC, colori=COLORI_DOC, conti=[x for x in docs if x.tipo == "conto"], n=len(lista), nxt="/documenti")


def _tre_cf(t, nome=False):
    import unicodedata
    t = re.sub(r"[^A-Z]", "", unicodedata.normalize("NFD", (t or "").upper()).encode("ascii", "ignore").decode())
    c = [x for x in t if x not in "AEIOU"]
    v = [x for x in t if x in "AEIOU"]
    if nome and len(c) >= 4:
        c = [c[0], c[2], c[3]]
    return "".join(c + v + ["X", "X", "X"])[:3]


def verifica_cf(cod, nome, cognome):
    """Controlla un codice fiscale (anche omocodico) rispetto a nome e cognome. Ritorna il messaggio d'errore, o '' se coerente."""
    if not nome.strip() or not cognome.strip():
        return "Per il codice fiscale inserisci nome e cognome."
    if not re.fullmatch(r"[A-Z]{6}[0-9LMNPQRSTUV]{2}[ABCDEHLMPRST][0-9LMNPQRSTUV]{2}[A-Z][0-9LMNPQRSTUV]{3}[A-Z]", cod):
        return "Il codice fiscale non è valido: deve avere 16 caratteri (6 lettere, 2 cifre, 1 lettera, 2 cifre, 1 lettera, 3 cifre, 1 lettera)."
    disp = {c: v for c, v in zip("0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ", [1, 0, 5, 7, 9, 13, 15, 17, 19, 21, 1, 0, 5, 7, 9, 13, 15, 17, 19, 21, 2, 4, 18, 20, 11, 3, 6, 8, 12, 14, 16, 10, 22, 25, 24, 23])}
    tot = sum(disp[c] if i % 2 == 0 else (int(c) if c.isdigit() else ord(c) - 65) for i, c in enumerate(cod[:15]))
    if chr(65 + tot % 26) != cod[15]:
        return "Il codice fiscale non è valido: l'ultima lettera di controllo non torna."
    if cod[:3] != _tre_cf(cognome):
        return "Il cognome non corrisponde al codice fiscale (le prime 3 lettere non coincidono)."
    if cod[3:6] != _tre_cf(nome, True):
        return "Il nome non corrisponde al codice fiscale (le lettere 4-6 non coincidono)."
    g = int(cod[9:11].translate(str.maketrans("LMNPQRSTUV", "0123456789")))
    if not (1 <= g <= 31 or 41 <= g <= 71):
        return "Il giorno di nascita nel codice fiscale non è valido."
    return ""


def iban_ok(v):
    v = "".join((v or "").split()).upper()
    if not re.fullmatch(r"[A-Z]{2}\d{2}[A-Z0-9]{11,30}", v) or (v.startswith("IT") and len(v) != 27):
        return False
    return int("".join(str(int(c, 36)) for c in v[4:] + v[:4])) % 97 == 1


def verifica_doc(f, tipo):
    """Controlli di coerenza sul modulo documenti (gli stessi della pagina). Ritorna il messaggio d'errore o ''."""
    nome, cogn = f.get("nome_p", "").strip(), f.get("cognome_p", "").strip()
    if not nome or not cogn:
        return "Inserisci nome e cognome."
    if not all(re.fullmatch(r"[^\W\d_][^\d_]*", x) and not re.search(r"[^\w' .\-]", x) for x in (nome, cogn)):
        return "Nome e cognome devono contenere solo lettere."
    num = f.get("numero", "").strip()
    if tipo == "conto":
        if not f.get("ente", "").strip():
            return "Indica la banca."
        if not iban_ok(f.get("iban", "")):
            return "L'IBAN non è valido."
    elif tipo in CARTE:
        if not f.get("ente", "").strip():
            return "Indica la banca."
        if not re.fullmatch(r"\d{12,19}", re.sub(r"[\s.\-]", "", num)):
            return "Il numero della carta deve avere da 12 a 19 cifre."
        if f.get("c_iban", "").strip() and f.get("nuovo_conto") == "1" and not iban_ok(f.get("c_iban", "")):
            return "L'IBAN del conto non è valido."
    elif tipo not in ("altro", "cf"):
        if len(num) < 4:
            return "Il numero del documento non sembra valido."
    return ""


@app.route("/documenti/salva", methods=["POST"])
@login_required
def documento_salva():
    uid = session["uid"]
    tipo = request.form.get("tipo", "altro")
    if tipo not in TIPI_DOC:
        tipo = "altro"
    ente = request.form.get("ente", "").strip()[:120]
    if tipo == "altro":
        titolo = request.form.get("titolo", "").strip()[:120] or TIPI_DOC[tipo]
    elif tipo in ("banc", "cc", "deb", "conto") and ente:
        titolo = f"{ente.upper()} - {TIPI_DOC[tipo]}"[:120]
    else:
        titolo = TIPI_DOC[tipo]
    err = verifica_doc(request.form, tipo)
    if err:
        flash(err)
        return redirect("/documenti")
    if tipo == "cf":
        ente = ""
        err = verifica_cf("".join(request.form.get("numero", "").split()).upper(), request.form.get("nome_p", ""), request.form.get("cognome_p", ""))
        if err:
            flash(err)
            return redirect("/documenti")
    d = None
    if request.form.get("id", "").isdigit():
        d = Documento.query.filter_by(id=int(request.form["id"]), user_id=uid).first()
    nuovo = d is None
    if nuovo:
        d = Documento(user_id=uid)
        db.session.add(d)
    d.tipo, d.titolo = tipo, titolo
    d.nome_p = request.form.get("nome_p", "").strip()[:60]
    d.cognome_p = request.form.get("cognome_p", "").strip()[:60]
    d.intestatario = f"{d.nome_p} {d.cognome_p}".strip()
    d.numero = request.form.get("numero", "").strip()[:60]
    if tipo == "cf":
        d.numero = "".join(d.numero.split()).upper()[:16]
    d.ente = ente
    d.rilascio = None if tipo in ("banc", "cc", "deb", "conto", "cf") else _data_doc(request.form.get("rilascio"))
    d.scadenza = None if tipo in ("conto", "cf") else _data_doc(request.form.get("scadenza"))
    d.iban = "".join(request.form.get("iban", "").split()).upper()[:40] if tipo == "conto" else ""
    d.nota = request.form.get("nota", "").strip()[:500]
    if tipo not in CARTE or request.form.get("rimuovi_img") == "1":
        d.immagine = d.img_tipo = None
    f = request.files.get("immagine")
    if f and tipo in CARTE:
        raw = f.read(700_001)
        mime = ("image/jpeg" if raw[:3] == b"\xff\xd8\xff" else "image/png" if raw[:8] == b"\x89PNG\r\n\x1a\n"
                else "image/webp" if raw[:4] == b"RIFF" and raw[8:12] == b"WEBP" else None)
        if raw and not mime:
            flash("Immagine non valida: usa un file JPG, PNG o WebP.")
        elif len(raw) > 700_000:
            flash("Immagine troppo grande (massimo 700 KB).")
        elif raw:
            d.immagine, d.img_tipo = base64.b64encode(raw).decode(), mime
    d.conto_id = None
    if tipo in CARTE:
        cid = request.form.get("conto_id", "")
        c = Documento.query.filter_by(id=int(cid), user_id=uid, tipo="conto").first() if cid.isdigit() else None
        iban_n = "".join(request.form.get("c_iban", "").split()).upper()[:40]
        c_int = request.form.get("c_intestatario", "").strip()[:120]
        c_num = request.form.get("c_numero", "").strip()[:60]
        if request.form.get("nuovo_conto") == "1" and (iban_n or c_int or c_num):
            c = Documento(user_id=uid, tipo="conto", titolo=(TIPI_DOC["conto"] + (" " + ente if ente else ""))[:120],
                          intestatario=c_int or d.intestatario, ente=ente, iban=iban_n, numero=c_num)
            db.session.add(c)
            db.session.flush()
        d.conto_id = c.id if c else None
    db.session.commit()
    flash("Documento inserito." if nuovo else "Documento aggiornato.")
    return redirect("/documenti")


@app.route("/documenti/ordine", methods=["POST"])
@login_required
def documenti_ordine():
    ids = [int(x) for x in request.form.get("ids", "").split(",") if x.isdigit()]
    mie = {d.id: d for d in Documento.query.filter(Documento.user_id == session["uid"], Documento.id.in_(ids))}
    for i, x in enumerate(ids):
        if x in mie:
            mie[x].ordine = i
    db.session.commit()
    return jsonify(ok=True)


@app.route("/documenti/<int:id>/immagine")
@login_required
def documento_immagine(id):
    d = Documento.query.filter_by(id=id, user_id=session["uid"]).first_or_404()
    if not d.immagine:
        abort(404)
    r = send_file(io.BytesIO(base64.b64decode(d.immagine)), mimetype=d.img_tipo or "image/jpeg")
    r.headers["Cache-Control"] = "private, max-age=86400"
    return r


@app.route("/documenti/<int:id>/elimina", methods=["POST"])
@login_required
def documento_elimina(id):
    d = Documento.query.filter_by(id=id, user_id=session["uid"]).first_or_404()
    Documento.query.filter_by(user_id=session["uid"], conto_id=d.id).update({"conto_id": None})
    db.session.delete(d)
    db.session.commit()
    flash("Documento eliminato.")
    return redirect("/documenti")


@app.route("/agenda")
@login_required
def agenda():
    vista = request.args.get("vista", "sett")
    if vista == "scad":
        return redirect("/scadenze")
    return _agenda("mese" if vista == "mese" else "sett")


@app.route("/scadenze")
@login_required
def scadenze():
    return _agenda("scad")


def _agenda(vista):
    uid = session["uid"]
    rif = _data(request.args.get("data")) or date.today()
    oggi = date.today()
    nomi = ["Lun", "Mar", "Mer", "Gio", "Ven", "Sab", "Dom"]
    ctxd = dict(colori=["#9333ea"] + COLORI_IMPEGNO, colori_sc=[""] + COLORI_IMPEGNO, vista=vista, oggi=oggi, nomi=nomi, mesi_brevi=[m[:3] for m in MESI_LUNGHI], data=rif.isoformat(),
                alert=alert_scadenze(uid))
    if vista == "sett":
        lun = rif - timedelta(days=rif.weekday())
        dom = lun + timedelta(days=6)
        ev = eventi_periodo(uid, lun, dom)
        giorni = [dict(d=lun + timedelta(days=i), ev=[], senza=[]) for i in range(7)]
        n_lez, ore_tot, n_imp = 0, 0.0, 0
        for g in giorni:
            for e in ev.get(g["d"], []):
                (g["ev"] if e["start"] is not None and e["kind"] != "scad" else g["senza"]).append(e)
                if e["kind"] == "lez":
                    n_lez += 1
                    ore_tot += e["ore"]
                elif e["kind"] == "imp":
                    n_imp += 1
            _colonne(g)
        inizi = [e["start"] // 60 for g in giorni for e in g["ev"]]
        fini = [-(-e["end"] // 60) for g in giorni for e in g["ev"]]
        h0, h1 = min([8] + inizi), min(max([21] + fini), 24)
        titolo = (f"{lun.day}–{dom.day} {MESI_LUNGHI[dom.month - 1]} {dom.year}" if lun.month == dom.month
                  else f"{lun.day} {MESI_LUNGHI[lun.month - 1]} – {dom.day} {MESI_LUNGHI[dom.month - 1]} {dom.year}")
        ctxd.update(giorni=giorni, h0=h0, h1=h1, titolo=titolo, n_sett=lun.isocalendar()[1], n_lez=n_lez, n_imp=n_imp, ore=ore_tot,
                    prec=(lun - timedelta(days=7)).isoformat(), succ=(lun + timedelta(days=7)).isoformat())
    elif vista == "mese":
        primo = rif.replace(day=1)
        ultimo = rif.replace(day=_cal.monthrange(rif.year, rif.month)[1])
        da = primo - timedelta(days=primo.weekday())
        a = ultimo + timedelta(days=6 - ultimo.weekday())
        ev = eventi_periodo(uid, da, a)
        settimane, d = [], da
        while d <= a:
            settimane.append([dict(d=d + timedelta(days=i), fuori=(d + timedelta(days=i)).month != rif.month,
                                   ev=ev.get(d + timedelta(days=i), [])) for i in range(7)])
            d += timedelta(days=7)
        ctxd.update(settimane=settimane, titolo=f"{MESI_LUNGHI[rif.month - 1].capitalize()} {rif.year}",
                    prec=aggiungi_mesi(primo, -1).isoformat(), succ=aggiungi_mesi(primo, 1).isoformat(),
                    nome_prec=MESI_LUNGHI[aggiungi_mesi(primo, -1).month - 1].capitalize(),
                    nome_succ=MESI_LUNGHI[aggiungi_mesi(primo, 1).month - 1].capitalize())
    else:
        voci = voci_scadenze(uid)
        tutti = Scadenza.query.filter_by(user_id=uid).all()
        per_anno = {}
        for s in tutti:
            if s.data_pagamento and s.importo:
                per_anno[s.data_pagamento.year] = per_anno.get(s.data_pagamento.year, 0) + s.importo
        note = sorted([s for s in tutti if s.tipo == "nota"], key=lambda s: (s.data_pagamento or date.min, s.id), reverse=True)
        attive = [v for v in voci if v["stato"] not in ("chiusa", "nd") and v["giorni"] is not None]
        prossima = min((v for v in attive if v["giorni"] >= 0), key=lambda v: v["giorni"], default=None)
        ctxd.update(prossima=prossima, n_attive=len(attive), pagato_anno=dict(per_anno).get(oggi.year, 0), anno_corr=oggi.year)
        ctxd.update(voci=voci, note=note, per_anno=sorted(per_anno.items(), reverse=True), suggerimenti=SUGGERIMENTI_SCADENZE)
    return render_template("agenda.html", **ctxd)


def _torna():
    nxt = request.form.get("next", "/agenda")
    return redirect(nxt if nxt.startswith(("/agenda", "/scadenze")) else "/agenda")


@app.route("/agenda/impegno/salva", methods=["POST"])
@login_required
def impegno_salva():
    uid = session["uid"]
    d = _data(request.form.get("data"))
    titolo = request.form.get("titolo", "").strip()[:150]
    if not d or not titolo:
        flash("Servono almeno data e titolo.")
        return _torna()
    i = None
    if request.form.get("id", "").isdigit():
        i = Impegno.query.filter_by(id=int(request.form["id"]), user_id=uid).first()
    if not i:
        i = Impegno(user_id=uid)
        db.session.add(i)
    i.data, i.titolo = d, titolo
    i.ora_inizio, i.ora_fine = ("", "") if request.form.get("tutto") else (_ora(request.form.get("ini")), _ora(request.form.get("fine")))
    col = request.form.get("colore", "")
    i.colore = col if col in COLORI_IMPEGNO else ""
    i.luogo, i.nota = request.form.get("luogo", "").strip()[:150], request.form.get("nota", "").strip()[:500]
    db.session.commit()
    return _torna()


@app.route("/agenda/impegno/<int:id>/elimina", methods=["POST"])
@login_required
def impegno_elimina(id):
    i = Impegno.query.filter_by(id=id, user_id=session["uid"]).first_or_404()
    db.session.delete(i)
    db.session.commit()
    return _torna()


@app.route("/agenda/scadenza/salva", methods=["POST"])
@login_required
def scadenza_salva():
    uid = session["uid"]
    voce = request.form.get("voce", "").strip()[:120]
    if not voce:
        flash("Indica di cosa si tratta (es. Assicurazione auto).")
        return _torna()
    imp = _num(request.form.get("importo", "").strip().replace(",", "."))
    dp = _data(request.form.get("data_pagamento"))
    val = _num(request.form.get("validita", "").strip().replace(",", "."))
    mesi = int(val * (12 if request.form.get("unita") == "anni" else 1)) if val and 0 < val <= 1200 else None
    sc = _data(request.form.get("scadenza"))
    nota_sola = request.form.get("tipo") == "nota"
    if nota_sola:  # solo nota/pagamento: nessuna validita' ne' scadenza
        mesi, sc = None, None
    if not sc and dp and mesi:
        sc = aggiungi_mesi(dp, mesi) - timedelta(days=1)  # scade il giorno prima dell'anniversario
    if not mesi and dp and sc and sc > dp:
        mesi = max(1, round(((sc - dp).days + 1) / 30.44))
    pre = _num(request.form.get("preavviso", "30"))
    s = None
    if request.form.get("id", "").isdigit():
        s = Scadenza.query.filter_by(id=int(request.form["id"]), user_id=uid).first()
    if not s:
        s = Scadenza(user_id=uid)
        db.session.add(s)
    s.voce, s.importo, s.data_pagamento, s.validita_mesi, s.scadenza = voce, imp, dp, mesi, sc
    s.preavviso = int(pre) if pre and 30 <= pre <= 365 else 30
    s.nota = request.form.get("nota", "").strip()[:500]
    s.chiusa = False
    col = request.form.get("colore", "")
    s.colore = col if col in COLORI_IMPEGNO else ""
    nuova = s.id is None
    s.tipo = "nota" if nota_sola else "scad"
    db.session.commit()
    if nota_sola:
        flash("Nota inserita." if nuova else "Nota aggiornata.")
    else:
        flash(("Scadenza inserita: " if nuova else "Scadenza aggiornata: ") + f"{s.voce}" + (f", scade il {_it(s.scadenza)}." if s.scadenza else "."))
    return _torna()


@app.route("/agenda/scadenza/<int:id>/elimina", methods=["POST"])
@login_required
def scadenza_elimina(id):
    s = Scadenza.query.filter_by(id=id, user_id=session["uid"]).first_or_404()
    db.session.delete(s)
    db.session.commit()
    return _torna()


@app.route("/agenda/scadenza/<int:id>/chiudi", methods=["POST"])
@login_required
def scadenza_chiudi(id):
    s = Scadenza.query.filter_by(id=id, user_id=session["uid"]).first_or_404()
    s.chiusa = not s.chiusa
    db.session.commit()
    return _torna()
